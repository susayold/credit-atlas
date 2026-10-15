"""Append the 79k-row state panel to the D02 workbook without re-importing all sheets."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "apps/excel/templates/CreditAtlas-D02-Data-Reference.xlsx"
DATA = ROOT / "docs/data/nmdb-performance-state-quarterly-2026q2.csv"


def add_data_sheet(book) -> None:
    if "NMDB states" in book.sheetnames:
        del book["NMDB states"]
    sheet = book.create_sheet("NMDB states", book.sheetnames.index("FHFA HPI state"))
    sheet.sheet_view.showGridLines = False
    headers = ["Source", "Frequency", "Geography level", "State code", "State / DC", "Mortgage market", "Quarter", "Year", "Quarter number", "Quarter-end month", "Suppressed", "Series ID", "Rate (%)"]
    sheet.append(headers)
    with DATA.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        for row in reader:
            rate = None if row["value_pct"] == "" else float(row["value_pct"])
            sheet.append([row["source"], row["frequency"], row["geo_level"], row["geo_id"], row["geo_name"],
                          row["market"], row["period"], int(row["year"]), int(row["quarter"]), int(row["month"]),
                          row["suppressed"] == "1", row["series_id"], rate])
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:M{sheet.max_row}"
    table = Table(displayName="NmdbStates", ref=f"A1:M{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    widths = [13, 13, 16, 12, 24, 35, 13, 10, 16, 19, 14, 13, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 13).number_format = '0.0"%"'
        sheet.cell(row, 11).number_format = "General"


def add_definitions(book) -> None:
    if "NMDB state definitions" in book.sheetnames:
        del book["NMDB state definitions"]
    sheet = book.create_sheet("NMDB state definitions", book.sheetnames.index("FHFA HPI definitions"))
    sheet.sheet_view.showGridLines = False
    rows = [
        ["Topic", "Definition / use", "Interpretation boundary"],
        ["Geography", "50 states plus District of Columbia; quarterly 2002Q1–2026Q2.", "State aggregates are not borrower residence rows or a property-level lookup."],
        ["P3089DL", "Active mortgage share reported 30 or 60 days past due at quarter end.", "An active-loan performance-state share, not fixed-horizon default probability."],
        ["P90DL", "Active mortgage share reported 90, 120, 150 or 180 days past due at quarter end.", "Keep the publisher's denominator and status definition."],
        ["PFORECL", "Active mortgage share in foreclosure, bankruptcy or deed-in-lieu status.", "A distinct performance category."],
        ["PFORB", "Active mortgage share reported in forbearance.", "Published from 2019Q4; prior blank cells mean unavailable, not zero."],
        ["Suppression", "FHFA suppresses small state cells with fewer than three loan records.", "68 values are suppressed in this snapshot; blank is not zero."],
        ["Preliminary periods", "The latest two quarters are preliminary due to credit-bureau reporting lag.", "2026Q1 and 2026Q2 may be revised."],
        ["History limitation", "Seriously past-due status may be undercounted from 2006 to 2012 due to credit-history purging.", "Do not attribute all period movement to borrower behavior."],
        ["Use boundary", "Aggregate 5% NMDB sample; use as state-level portfolio context and trend comparison.", "Do not train D01, infer individual borrower PD, or join to Freddie/NSMO loans."],
        ["Source", "FHFA NMDB mortgage-performance statistics — States.", "https://www.fhfa.gov/data/nmdb"],
    ]
    for values in rows:
        sheet.append(values)
    sheet.freeze_panes = "B2"
    table = Table(displayName="NmdbStateDefinitions", ref=f"A1:C{len(rows)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col, width in {"A": 24, "B": 90, "C": 92}.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    book = load_workbook(WORKBOOK)
    add_data_sheet(book)
    add_definitions(book)
    guide = book["Start here"]
    guide.append(["NMDB states", "State/DC quarterly aggregate panel through 2026Q2; filter by state, market, series and quarter. Suppressed or unavailable values remain blank."])
    last = guide.max_row
    for cell in guide[last]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.cell(last, 1).font = Font(bold=True, color="243B53")
    guide.row_dimensions[last].height = 38
    book.save(WORKBOOK)
    print(f"Added NMDB states ({book['NMDB states'].max_row - 1:,} rows) and definitions to {WORKBOOK}")


if __name__ == "__main__":
    main()
