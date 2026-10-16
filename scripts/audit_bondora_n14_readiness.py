"""Aggregate-only Bondora label maturity and country/time support audit."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHEETS = ["Loan Dataset", "Dataset Dictionary"]
REQUIRED_COLUMNS = {
    "loan_id", "country", "loan_issued_at", "loan_last_recorded_action_date_local",
    "has_default_within_12_months", "loan_status", "is_default", "months_on_book",
}
SMALL_CELL_THRESHOLD = 50


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def label_class(value) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)) and value in (0, 1):
        return "true" if value == 1 else "false"
    return "other"


def add_twelve_calendar_months(value: date) -> date:
    """Return the same calendar day next year, clamped at month end for leap day."""
    try:
        return value.replace(year=value.year + 1)
    except ValueError:
        return value.replace(year=value.year + 1, day=28)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path, help="Private official Bondora XLSX")
    parser.add_argument("--out", required=True, type=Path, help="Aggregate-only JSON output path (keep private until rights review)")
    parser.add_argument("--snapshot-proxy", required=True, type=date.fromisoformat,
                        help="Publisher file Last-Modified date, not a verified observation cutoff")
    parser.add_argument("--expected-sha256", help="Optional pinned source SHA-256 from the acquisition receipt")
    args = parser.parse_args()
    source = args.source.resolve()
    if not source.is_file():
        raise SystemExit(f"Missing source workbook: {source}")
    source_hash = sha256(source)
    if args.expected_sha256 and source_hash.lower() != args.expected_sha256.lower():
        raise SystemExit("Source hash differs from the expected acquisition receipt")

    workbook = load_workbook(source, read_only=True, data_only=True)
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Loan Dataset"]
    header = [str(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if len(header) != 31 or len(set(header)) != len(header) or not REQUIRED_COLUMNS.issubset(header):
        raise ValueError("Unexpected source column contract")
    idx = {name: i for i, name in enumerate(header)}

    by_year = defaultdict(Counter)
    by_country = defaultdict(Counter)
    by_country_year = defaultdict(Counter)
    maturity_by_year = defaultdict(Counter)
    maturity_by_country = defaultdict(Counter)
    maturity_by_country_year = defaultdict(Counter)
    mature_null_diagnostics = {
        "loan_status": Counter(), "is_default": Counter(),
        "action_observation": Counter(), "months_on_book": Counter(),
        "country_year": Counter(),
    }
    country_dates = defaultdict(lambda: {"min": None, "max": None})
    all_labels = Counter()
    unexpected_labels = Counter()
    unique_ids: set[object] = set()
    duplicate_ids = 0
    issue_date_missing = 0
    latest_action = None
    maturity = Counter()
    row_count = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row is None or all(value is None for value in row):
            continue
        row_count += 1
        loan_id = row[idx["loan_id"]]
        if loan_id in unique_ids:
            duplicate_ids += 1
        else:
            unique_ids.add(loan_id)
        label = label_class(row[idx["has_default_within_12_months"]])
        all_labels[label] += 1
        if label == "other":
            unexpected_labels[type(row[idx["has_default_within_12_months"]]).__name__] += 1

        issued = as_date(row[idx["loan_issued_at"]])
        if issued is None:
            issue_date_missing += 1
            continue
        year = str(issued.year)
        country_value = row[idx["country"]]
        country = str(country_value).strip() if country_value is not None else "<NULL>"
        by_year[year][label] += 1
        by_country[country][label] += 1
        by_country_year[(country, year)][label] += 1
        bounds = country_dates[country]
        bounds["min"] = issued if bounds["min"] is None or issued < bounds["min"] else bounds["min"]
        bounds["max"] = issued if bounds["max"] is None or issued > bounds["max"] else bounds["max"]

        action = as_date(row[idx["loan_last_recorded_action_date_local"]])
        if action is not None:
            latest_action = action if latest_action is None or action > latest_action else latest_action

        mature_by_proxy = add_twelve_calendar_months(issued) <= args.snapshot_proxy
        maturity_class = "mature_by_file_date_proxy" if mature_by_proxy else "not_mature_by_proxy"
        maturity[(maturity_class, label)] += 1
        maturity_by_year[(year, maturity_class)][label] += 1
        maturity_by_country[(country, maturity_class)][label] += 1
        maturity_by_country_year[(country, year, maturity_class)][label] += 1
        if mature_by_proxy and label == "null":
            status = row[idx["loan_status"]]
            mature_null_diagnostics["loan_status"][str(status) if status is not None else "<NULL>"] += 1
            default_value = row[idx["is_default"]]
            mature_null_diagnostics["is_default"][label_class(default_value)] += 1
            if action is None:
                mature_null_diagnostics["action_observation"]["missing"] += 1
            elif action < add_twelve_calendar_months(issued):
                mature_null_diagnostics["action_observation"]["last_action_before_12m_anniversary"] += 1
            else:
                mature_null_diagnostics["action_observation"]["last_action_on_or_after_12m_anniversary"] += 1
            months_on_book = row[idx["months_on_book"]]
            if months_on_book is None:
                mob_group = "missing"
            elif isinstance(months_on_book, (int, float)) and months_on_book < 12:
                mob_group = "under_12"
            elif isinstance(months_on_book, (int, float)):
                mob_group = "12_or_more"
            else:
                mob_group = "other"
            mature_null_diagnostics["months_on_book"][mob_group] += 1
            mature_null_diagnostics["country_year"][f"{country}/{year}"] += 1

        if row_count % 100_000 == 0:
            print(f"audited_rows={row_count}", flush=True)

    if row_count != 789_144 or duplicate_ids or issue_date_missing:
        raise ValueError(f"Unexpected source integrity: rows={row_count}, duplicate_ids={duplicate_ids}, missing_issue_dates={issue_date_missing}")
    if set(all_labels) - {"true", "false", "null"}:
        raise ValueError(f"Unexpected target value types: {dict(unexpected_labels)}")

    year_report = {
        year: {"loans": sum(counts.values()), "target_true": counts["true"],
               "target_false": counts["false"], "target_null": counts["null"]}
        for year, counts in sorted(by_year.items())
    }
    country_report = {}
    for country, counts in sorted(by_country.items()):
        date_range = country_dates[country]
        country_report[country] = {
            "loans": sum(counts.values()), "target_true": counts["true"],
            "target_false": counts["false"], "target_null": counts["null"],
            "issue_date_min": date_range["min"].isoformat(),
            "issue_date_max": date_range["max"].isoformat(),
        }
    country_year_report = []
    for (country, year), counts in sorted(by_country_year.items()):
        total = sum(counts.values())
        suppressed = total < SMALL_CELL_THRESHOLD
        country_year_report.append({
            "country": country, "issue_year": year, "loans": None if suppressed else total,
            "target_true": None if suppressed else counts["true"],
            "target_false": None if suppressed else counts["false"],
            "target_null": None if suppressed else counts["null"],
            "small_cell_suppressed": suppressed,
        })

    maturity_report = {
        "rule": "loan_issued_at + 12 calendar months <= publisher file Last-Modified date",
        "snapshot_proxy": args.snapshot_proxy.isoformat(),
        "latest_recorded_action_date_in_rows": latest_action.isoformat() if latest_action else None,
        "counts": {f"{maturity_class}_{label}": count
                   for (maturity_class, label), count in sorted(maturity.items())},
        "by_origination_year": [
            {"issue_year": year, "maturity_group": group,
             "loans": sum(counts.values()), "target_true": counts["true"],
             "target_false": counts["false"], "target_null": counts["null"]}
            for (year, group), counts in sorted(maturity_by_year.items())
        ],
        "by_country": [
            {"country": country, "maturity_group": group,
             "loans": sum(counts.values()), "target_true": counts["true"],
             "target_false": counts["false"], "target_null": counts["null"]}
            for (country, group), counts in sorted(maturity_by_country.items())
        ],
        "by_country_and_year": [
            {"country": country, "issue_year": year, "maturity_group": group,
             "loans": None if sum(counts.values()) < SMALL_CELL_THRESHOLD else sum(counts.values()),
             "target_true": None if sum(counts.values()) < SMALL_CELL_THRESHOLD else counts["true"],
             "target_false": None if sum(counts.values()) < SMALL_CELL_THRESHOLD else counts["false"],
             "target_null": None if sum(counts.values()) < SMALL_CELL_THRESHOLD else counts["null"],
             "small_cell_suppressed": sum(counts.values()) < SMALL_CELL_THRESHOLD}
            for (country, year, group), counts in sorted(maturity_by_country_year.items())
        ],
        "proxy_mature_null_diagnostic_only": {
            name: dict(sorted(counts.items()))
            for name, counts in mature_null_diagnostics.items()
        },
        "interpretation": "Diagnostic proxy only. File Last-Modified is not a verified per-loan observation cutoff; mismatches do not prove a bad source label or a measured negative.",
    }
    report = {
        "status": "PASS_AGGREGATE_ONLY_DIAGNOSTIC",
        "source": {"sha256": source_hash, "bytes": source.stat().st_size,
                   "sheets": workbook.sheetnames, "data_rows": row_count, "columns": len(header)},
        "row_level_ids_emitted": False,
        "model_fit_performed": False,
        "label_definition_from_source_dictionary": "Boolean: default observed in 12 months on book; NULL when not measurable yet.",
        "label_counts": {"true": all_labels["true"], "false": all_labels["false"], "null": all_labels["null"]},
        "by_origination_year": year_report,
        "by_country": country_report,
        "by_country_and_year": country_year_report,
        "maturity_proxy_diagnostic": maturity_report,
        "feature_use_control": {
            "application_time_feature_set_created": False,
            "post_booking_columns_used_as_predictors": False,
            "availability_note": "The source dictionary does not establish point-in-time availability for all issue-date and income fields. No predictors are selected or fitted by this audit.",
        },
        "rights_note": "Aggregate output only. The public download does not itself establish row-level redistribution or commercial reuse rights; keep source workbook and detailed output private pending review.",
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    workbook.close()
    print(json.dumps({"status": report["status"], "rows": row_count, "labels": report["label_counts"],
                      "countries": len(country_report), "country_year_cells": len(country_year_report),
                      "maturity_proxy": maturity_report["counts"], "output": str(args.out)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
