from __future__ import annotations

import hashlib
import json
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/private/retail/bondora-source/loan_dataset_investor.xlsx"
OUT = ROOT / "artifacts/private/bondora-public-dataset-profile"


def serial(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def main():
    if not SOURCE.is_file():
        raise SystemExit(f"Missing source workbook: {SOURCE}")
    OUT.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    with SOURCE.open("rb") as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(block)

    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    if workbook.sheetnames != ["Loan Dataset", "Dataset Dictionary"]:
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    dictionary = {}
    for row in workbook["Dataset Dictionary"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            dictionary[str(row[0])] = str(row[1] or "")

    sheet = workbook["Loan Dataset"]
    header = [str(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if len(header) != 31 or len(set(header)) != len(header):
        raise ValueError(f"Unexpected header count/duplicates: {len(header)}")
    index = {name: pos for pos, name in enumerate(header)}
    required = {
        "loan_id", "country", "loan_issued_at", "loan_last_recorded_action_date_local",
        "loan_status", "is_default", "has_default_within_12_months", "months_on_book",
    }
    if not required.issubset(index):
        raise ValueError(f"Missing required columns: {sorted(required - set(index))}")

    nulls = Counter()
    countries = Counter()
    statuses = Counter()
    target = Counter()
    years = defaultdict(lambda: {"loans": 0, "target_true": 0, "target_false": 0, "target_null": 0})
    unique_ids = set()
    duplicate_ids = 0
    issue_min = issue_max = action_max = None
    row_count = 0
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if values is None or all(value is None for value in values):
            continue
        row_count += 1
        for name, pos in index.items():
            if values[pos] is None:
                nulls[name] += 1
        loan_id = values[index["loan_id"]]
        if loan_id in unique_ids:
            duplicate_ids += 1
        else:
            unique_ids.add(loan_id)
        country = values[index["country"]]
        countries[str(country) if country is not None else "<NULL>"] += 1
        status = values[index["loan_status"]]
        statuses[str(status) if status is not None else "<NULL>"] += 1
        label = values[index["has_default_within_12_months"]]
        if label is None:
            label_key = "null"
        elif label is True or label == 1:
            label_key = "true"
        elif label is False or label == 0:
            label_key = "false"
        else:
            label_key = f"other:{label}"
        target[label_key] += 1
        issued = values[index["loan_issued_at"]]
        if isinstance(issued, (datetime, date)):
            issue_min = issued if issue_min is None or issued < issue_min else issue_min
            issue_max = issued if issue_max is None or issued > issue_max else issue_max
            year = str(issued.year)
            years[year]["loans"] += 1
            if label_key in {"true", "false", "null"}:
                years[year][f"target_{label_key}"] += 1
        action = values[index["loan_last_recorded_action_date_local"]]
        if isinstance(action, (datetime, date)):
            action_max = action if action_max is None or action > action_max else action_max
        if row_count % 100_000 == 0:
            print(f"profiled_rows={row_count}", flush=True)

    missing_dictionary = sorted(set(header) - set(dictionary))
    profile = {
        "source": {
            "publisher": "Go & Grow (formerly Bondora), public statistics page",
            "download_url": "https://sabanners001.blob.core.windows.net/statistics/public/loan_dataset_investor.xlsx",
            "landing_page": "https://goandgrow.eu/en/public-statistics/",
            "file": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
            "bytes": SOURCE.stat().st_size,
            "sha256": digest.hexdigest(),
            "sheets": workbook.sheetnames,
        },
        "table": {
            "data_rows": row_count,
            "columns": len(header),
            "headers": header,
            "unique_loan_ids": len(unique_ids),
            "duplicate_loan_id_rows": duplicate_ids,
            "issue_date_min": serial(issue_min),
            "issue_date_max": serial(issue_max),
            "last_recorded_action_max": serial(action_max),
            "country_counts": dict(sorted(countries.items())),
            "loan_status_counts": dict(sorted(statuses.items())),
            "has_default_within_12_months_counts": dict(sorted(target.items())),
            "origination_year_label_counts": dict(sorted(years.items())),
            "null_counts": {name: nulls[name] for name in header},
        },
        "dictionary": {
            "entries": len(dictionary),
            "missing_columns": missing_dictionary,
            "definitions": dictionary,
        },
        "interpretation": {
            "supports_application_underwriting": False,
            "reason": "The public workbook is a loan-level issued-loan/investor dataset. Its published schema does not include application decision date, requested amount, decision outcome/rejected applications, or bureau fields. Combined income has no documented approval-time availability in this workbook.",
            "supports_bounded_12_month_default_analysis": True,
            "label_caveat": "The source dictionary defines has_default_within_12_months as a Boolean with NULL when not measurable yet. Preserve NULL as immature/unobserved, not as non-default; independently verify snapshot timing and target construction before training.",
            "privacy_caveat": "Keep loan identifiers and row-level data in private research storage; expose only aggregated, qualified outputs.",
            "reuse_caveat": "The publisher exposes a download link, but this audit does not establish permission for redistribution or commercial reuse. Review applicable terms and request permission before publishing the workbook or row-level extracts.",
        },
    }
    out = OUT / "profile.json"
    out.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    receipt = {
        "status": "PROFILED",
        "audit_script": str(Path(__file__).relative_to(ROOT)).replace("\\", "/"),
        "audit_script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "profile": str(out.relative_to(ROOT)).replace("\\", "/"),
        "profile_sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
        "source_sha256": digest.hexdigest(),
        "data_rows": row_count,
        "columns": len(header),
        "unique_loan_ids": len(unique_ids),
        "duplicate_loan_id_rows": duplicate_ids,
        "pii_or_row_ids_emitted": False,
    }
    receipt_path = OUT / "receipt.json"
    receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "receipt": str(receipt_path.relative_to(ROOT)).replace("\\", "/"),
        "rows": row_count,
        "columns": len(header),
        "unique_ids": len(unique_ids),
        "duplicate_id_rows": duplicate_ids,
        "issue_range": [serial(issue_min), serial(issue_max)],
        "action_max": serial(action_max),
        "countries": countries,
        "statuses": statuses,
        "target": target,
        "year_target_counts": years,
    }, ensure_ascii=False, default=dict))


if __name__ == "__main__":
    main()
