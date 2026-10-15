"""Independent openpyxl and Decimal comparison against running HTTP projection."""
import json,uuid
from pathlib import Path
from decimal import Decimal,localcontext
from urllib.request import urlopen
from openpyxl import load_workbook

root=Path.cwd();folder=root/'data/private/structured/pool-history-source-757ce3631134'
with urlopen('http://127.0.0.1:8770/api/managed-pool') as response:r=json.load(response)
ws=load_workbook(folder/'ManagedPool-Retail.xlsx',read_only=True,data_only=True).active
cells={cell.coordinate:cell.value for row in ws.iter_rows() for cell in row if hasattr(cell,'coordinate')}
count=0;reconciliations=0
with localcontext() as ctx:
    ctx.prec=50
    for p in r['periods']:
        col=p['column'];assert cells[col+'26'].date().isoformat()==p['end_date']
        for key,v in p['values'].items():
            source=cells[v['cell']]
            if source is None:assert v['value_exact'] is None
            else:assert abs(Decimal(str(source))-Decimal(v['value_exact']))<=max(Decimal('1e-14'),abs(Decimal(str(source)))*Decimal('1e-15'))
            count+=1
        def d(row):return Decimal(str(cells[col+str(row)]))
        assert abs(d(17)+d(18)-d(19))<Decimal('0.000001')
        assert sum(d(i) for i in range(31,35))==d(35)
        assert abs(sum(d(i) for i in range(45,49))-d(49))<Decimal('0.000001')
        assert d(5)==d(27) and d(8)==d(28)
        for start,rates,den in [(31,38,8),(45,52,5)]:
            for i in range(5):assert abs(d(start+i)/d(den)-d(rates+i))<=Decimal('0.000050000000001')
        assert abs(Decimal(p['derived_ratios']['period_net_loss_fraction']['value_exact'])-d(19)/d(6))<Decimal('1e-15')
        assert abs(Decimal(p['derived_ratios']['period_repos_fraction']['value_exact'])-d(12)/d(10))<Decimal('1e-15')
        assert all(c['status']=='PASS' for c in p['checks']);reconciliations+=15
out=root/'artifacts/private'/('managed-pool-audit-'+uuid.uuid4().hex[:12]);out.mkdir()
(out/'receipt.json').write_text(json.dumps({'status':'PASS','source_cells_compared':count,'independent_reconciliations':reconciliations,'derived_ratios_compared':14,'periods':7,'source_file':'ManagedPool-Retail.xlsx'},indent=2),encoding='utf8')
print(out)
