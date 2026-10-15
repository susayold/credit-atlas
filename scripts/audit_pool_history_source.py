"""Read-only workbook audit with an independent Excel reader."""
from pathlib import Path
import openpyxl,json,uuid
from decimal import Decimal
from urllib.request import urlopen
root=Path.cwd();base='http://127.0.0.1:8770/api/pool-history'
with urlopen(base) as response:report=json.load(response)
folder=root/'data/private/structured'/report['source_id'];book=openpyxl.load_workbook(folder/'StaticPool-WOART.xlsx',data_only=True,read_only=True)
count=missing=0;checks=[]
for metric,sheet in [('delinquency','Delinquencies'),('loss','Cumulative Losses'),('speed','ABS Speed')]:
    s=book[sheet];cached=list(s.iter_rows())
    for column in range(2,23):
        deal=cached[2][column-1].value
        with urlopen(base+'?metric='+metric+'&deal='+deal.replace(' ','%20')) as response:r=json.load(response)
        by_age={x['collection_age']:x for x in r['rows']}
        for row in range(4,s.max_row+1):
            age=cached[row-1][0].value
            if not isinstance(age,(int,float)):continue
            source=cached[row-1][column-1];projected=by_age[int(age)]
            if source.value is not None:assert projected['cell']==source.coordinate
            if source.value is None:assert projected['value_exact'] is None;missing+=1
            else:
                assert abs(Decimal(str(source.value))-Decimal(projected['value_exact']))<=Decimal('1e-15')
                count+=1
        checks.append({'metric':metric,'deal':deal,'status':'PASS'})
assert count==report['inventory']['populated_age_cells']==1998
assert missing==report['inventory']['blank_age_cells']==1362
out=root/'artifacts/private'/('pool-history-audit-'+uuid.uuid4().hex[:12]);out.mkdir()
(out/'receipt.json').write_text(json.dumps({'status':'PASS','populated_cells':count,'missing_cells':missing,'cohort_metric_groups':len(checks),'checks':checks,'reader':'openpyxl readonly cache reader; API module not imported','formula_recalculated':False},indent=2),encoding='utf8')
print(out);print(count,missing)
