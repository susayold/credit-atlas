"""Independent openpyxl audit of every normalized observation and source cell."""
import argparse,hashlib,json,uuid
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

root=Path.cwd()
parser=argparse.ArgumentParser();parser.add_argument('--profile',default='artifacts/private/sovereign-default-profile-08270bdda389/profile.json');args=parser.parse_args()
profile=(root/args.profile).resolve()
if not profile.is_relative_to(root/'artifacts/private'):raise ValueError('Private profile required')
doc=json.loads(profile.read_text(encoding='utf8'))
source=root/'data/private/country'/doc['source_id']
receipt=json.loads((source/'acquisition-complete.json').read_text(encoding='utf8'))
hashes={s['file']:hashlib.sha256((source/s['file']).read_bytes()).hexdigest()==s['sha256'] for s in receipt['sources']}
if doc.get('archive_source_id'):
    archive=root/'data/private/country'/doc['archive_source_id']
    archive_receipt=json.loads((archive/'acquisition-complete.json').read_text(encoding='utf8'))
    hashes.update({s['file']:hashlib.sha256((archive/s['file']).read_bytes()).hexdigest()==s['sha256'] for s in archive_receipt['sources']})
issues=[];checked=0;missing=0;zero=0;markers=0;absent=0;edition_checks=[]
for edition in doc['editions']:
    edition_source=archive if edition.get('source_id')==doc.get('archive_source_id') and doc.get('archive_source_id') else source
    wb=load_workbook(edition_source/edition['file'],read_only=True,data_only=True,keep_links=False)
    summary,flat=wb.worksheets
    it=flat.iter_rows(values_only=True)
    header_row=edition.get('header_row',65)
    for _ in range(header_row-1):next(it)
    headers=next(it);indices={v:i for i,v in enumerate(headers) if v is not None}
    observations={}
    for n,row in enumerate(it,header_row+1):
        key=row[indices['k3' if edition['edition']==2026 else 'k']]
        if key:observations[key]=(n,row)
    entities={e['id']:e for e in edition['entities']}
    for record in edition['records']:
        key=f"{record['entity']}_{record['year']}"
        n,row=observations.pop(key)
        if n!=record['source_row'] or row[indices['DEBT_YEAR']]!=record['year']:issues.append([key,'row/year'])
        if row[indices['DEBT_COUNTRY']]!=entities[record['entity']]['name']:issues.append([key,'name'])
        if row[indices['DEBT_COUNTRY_GROUP']]!=entities[record['entity']]['source_group']:issues.append([key,'group'])
        for metric,actual in record['values'].items():
            source_header=edition.get('metric_source_headers',{}).get(metric,f"DEBT_{metric}_{edition['edition']}")
            checked+=1
            if source_header is None:
                absent+=1
                if actual is not None:issues.append([key,metric,'absent schema became numeric'])
                continue
            expected=row[indices[source_header]]
            if isinstance(expected,str):
                markers+=1;marker=record.get('raw_non_numeric_values',{}).get(metric)
                if actual is not None or not marker or marker['raw']!=expected:issues.append([key,metric,'source marker lost'])
                continue
            if expected is None:
                missing+=1
                if actual is not None:issues.append([key,metric,'missing became numeric'])
            elif actual is None:issues.append([key,metric,'numeric became missing'])
            else:
                # openpyxl converts source XML to binary floats. The production
                # parser preserves XML text; permit only that conversion error.
                x=Decimal(str(expected));a=Decimal(actual)
                if abs(a-x)>max(Decimal('1e-9'),abs(x)*Decimal('1e-14')):issues.append([key,metric,'value mismatch'])
                if x==0:zero+=1
    if observations:issues.append([edition['edition'],'unparsed rows',len(observations)])
    summary_headers={str(c.value):c.coordinate for c in next(summary.iter_rows(min_row=1,max_row=1)) if c.value is not None}
    for h in edition['year_headers']:
        if summary_headers.get(h['source_label'])!=h['source_cell']:issues.append([h['year'],'header'])
    conflicts=[c for c in edition['source_summary_comparisons'] if not c['matches_at_flat_precision']]
    edition_checks.append({'edition':edition['edition'],'rows':len(edition['records']),'entities':len(entities),
                          'summary_conflicts':conflicts,'latest_year_header':edition['year_headers'][-1]})
    wb.close()
out=root/'artifacts/private'/('sovereign-default-audit-'+uuid.uuid4().hex[:12]);out.mkdir()
result={'status':'PASS' if not issues and all(hashes.values()) else 'ISSUES_FOUND',
        'profile_sha256':hashlib.sha256(profile.read_bytes()).hexdigest(),'source_hash_checks':hashes,
        'values_checked':checked,'missing_values_checked':missing,'zero_values_checked':zero,
        'source_marker_values_checked':markers,'schema_absent_values_checked':absent,
        'issues':issues,'editions':edition_checks,'scope':'Every selected flat-sheet metric, source row, entity metadata and year header; XML float precision preserved separately',
        'source_summary_conflicts_resolved':False,'pd_lgd_calibrated':False,'macro_stress_linkage_qualified':False}
(out/'receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(out);print(result['status']);print(checked,missing,zero);print(issues[:10])
if result['status']!='PASS':raise SystemExit(1)
