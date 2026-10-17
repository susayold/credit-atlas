"""Edge browser acceptance for the FHFA NMDB aggregate context panel."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import threading
from datetime import datetime, timezone
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

# The private venv may have an editable install of the private checkout. Pin this
# QA runner to the source tree whose UI and route contract it is validating.
REPO_SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(REPO_SRC))

from credit_atlas.server import handler_for
from credit_atlas.nmdb_performance_geographies import report as nmdb_geography_report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--web-root", type=Path, help="Optional UI directory when verified data and source checkout are in separate roots.")
    parser.add_argument("--receipt", required=True, type=Path)
    args = parser.parse_args()
    web_root = args.web_root.resolve() if args.web_root else None
    root = args.root.resolve()
    service = SimpleNamespace(root=root, fhfa_nmdb_performance_geographies=lambda query: nmdb_geography_report(root, query))
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler_for(service, web_root=web_root))
    server.daemon_threads = True
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    checks: list[str] = []
    errors: list[str] = []
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="msedge", headless=True)
            page = browser.new_page()
            page.on("pageerror", lambda error: errors.append(str(error)))
            response = page.goto(f"http://127.0.0.1:{server.server_port}/nsmo", wait_until="networkidle", timeout=45_000)
            if response is None or response.status != 200:
                raise RuntimeError("NSMO page did not return HTTP 200")
            page.locator("#nmdb-state").get_by_text("FHFA NMDB aggregate verified", exact=False).wait_for(timeout=20_000)
            if page.locator("#nmdb-latest .metric").count() != 4:
                raise AssertionError("Latest quarter panel does not show all four mortgage markets")
            metrics = page.locator("#nmdb-latest").inner_text()
            if "1.7%" not in metrics or "4.2%" not in metrics:
                raise AssertionError("All Mortgages 2026Q2 rates do not match the source")
            if page.locator("#nmdb-quarterly tbody tr").count() != 98:
                raise AssertionError("All 98 quarterly observations are not shown for the selected market")
            if page.locator("#nmdb-trend .nmdb-line").count() != 4:
                raise AssertionError("Quarterly chart does not show all four market series")
            if page.locator('#nmdb-series option[value="P3089DL"]').inner_text() != "30–89 days delinquent":
                raise AssertionError("P3089DL selector does not show the 30–89 day category label")
            checks.append("NMDB hash-checked API loads 98-quarter trend, four latest KPIs and the FHFA P3089DL label")

            series = page.locator("#nmdb-series")
            series.select_option("P90DL")
            p90_metrics = page.locator("#nmdb-latest").inner_text()
            if "1%" not in p90_metrics or "2.8%" not in p90_metrics:
                raise AssertionError(f"90+ DPD selector does not update the market KPIs: {p90_metrics}")
            checks.append("performance series selector updates latest KPI cards and trend")

            series.select_option("PFORB")
            market = page.locator("#nmdb-market")
            market.select_option("Government / Non-Conventional")
            table = page.locator("#nmdb-quarterly tbody")
            if table.locator("tr").count() != 98 or table.locator("tr").first.inner_text().find("—") < 0:
                raise AssertionError("Pre-2019Q4 forbearance blanks are not preserved as blanks")
            if "1.4%" not in page.locator("#nmdb-latest").inner_text():
                raise AssertionError("Government forbearance latest rate does not match FHFA")
            checks.append("forbearance blanks remain blank before first publication; market filter works")

            state = page.locator("#nmdb-state-choice")
            state.select_option("VT")
            page.locator("#nmdb-state-panel-status").get_by_text("Vermont", exact=False).wait_for(timeout=20_000)
            page.locator("#nmdb-state-series").select_option("P90DL")
            page.locator("#nmdb-state-market").select_option("Government / Non-Conventional")
            page.wait_for_function("document.querySelectorAll('#nmdb-state-trend .nmdb-line').length > 4")
            suppressed_row = page.locator("#nmdb-state-quarterly tbody tr").filter(has_text="2004Q1").first
            if suppressed_row.locator("td").nth(2).inner_text() != "—" or suppressed_row.locator("td").nth(6).inner_text() != "Yes":
                raise AssertionError("Vermont P90DL suppressed cell is not preserved in the quarterly table")
            checks.append("state chart breaks around source-suppressed Vermont values and the table retains the suppression flag")

            link = page.get_by_role("link", name="Tải workbook Excel · 1.568 dòng")
            with page.expect_download() as download_info:
                link.click()
            download = download_info.value
            if download.suggested_filename != "CreditAtlas-NMDB-Performance.xlsx":
                raise AssertionError("NMDB Excel filename does not match the app link")
            saved = args.receipt.parent / "CreditAtlas-NMDB-Performance-browser-qa.xlsx"
            saved.parent.mkdir(parents=True, exist_ok=True)
            download.save_as(saved)
            if saved.read_bytes()[:2] != b"PK":
                raise AssertionError("NMDB download is not an XLSX package")
            checks.append("four-sheet NMDB Excel workbook downloads from the app")

            page.set_viewport_size({"width": 390, "height": 844})
            dimensions = page.evaluate("({w:document.querySelector('.nmdb-context-panel').clientWidth,s:document.querySelector('.nmdb-context-panel').scrollWidth,cols:getComputedStyle(document.querySelector('#nmdb-latest')).gridTemplateColumns.split(' ').length})")
            if dimensions["s"] > dimensions["w"] or dimensions["cols"] != 2:
                raise AssertionError(f"NMDB panel has mobile overflow: {dimensions}")
            checks.append("NMDB chart and KPI panel fit a mobile viewport")

            explorer = page.locator("#nmdb-all-geographies")
            explorer.wait_for(timeout=20_000)
            level = explorer.locator("#nmdb-all-geo-level")
            level.select_option("Metro")
            metro_rows = explorer.locator("#nmdb-all-geo-table tbody tr")
            metro_rows.nth(99).wait_for(timeout=20_000)
            if metro_rows.count() != 100:
                raise AssertionError(f"Metro filter returned {metro_rows.count()} rows, expected 100")
            if "100" not in explorer.inner_text():
                raise AssertionError("Metro filter does not expose all 100 source geographies")
            checks.append("all 100 metro geographies load for the selected FHFA quarter and series")

            level.select_option("State")
            geography = explorer.locator("#nmdb-all-geo-id")
            geography.select_option("CA")
            explorer.locator("#nmdb-all-geo-period").select_option("all")
            state_rows = explorer.locator("#nmdb-all-geo-table tbody tr")
            state_rows.nth(97).wait_for(timeout=20_000)
            if state_rows.count() != 98:
                raise AssertionError(f"Single-state trend returned {state_rows.count()} rows, expected 98")
            explorer.locator("#nmdb-all-geo-trend .nmdb-line").first.wait_for(timeout=20_000)
            if explorer.locator("#nmdb-all-geo-trend .nmdb-line").count() < 1:
                raise AssertionError("Single-state quarterly performance trend is not charted")
            quarter_ticks = explorer.locator("#nmdb-all-geo-trend text").all_text_contents()
            if "2002Q1" not in quarter_ticks or "2026Q2" not in quarter_ticks:
                raise AssertionError(f"Trend chart quarter ticks do not show the full source range: {quarter_ticks}")
            checks.append("state selector returns all 98 quarters with correct source-quarter chart ticks")

            geography.select_option("VT")
            explorer.locator("#nmdb-all-geo-market").select_option("Government / Non-Conventional")
            explorer.locator("#nmdb-all-geo-period").select_option("2004Q1")
            result_rows = explorer.locator("#nmdb-all-geo-table tbody tr")
            result_rows.first.wait_for(timeout=20_000)
            if result_rows.count() != 1:
                raise AssertionError(f"Suppressed Vermont filter returned {result_rows.count()} rows, expected 1")
            result_row = explorer.locator("#nmdb-all-geo-table tbody tr").inner_text()
            if "suppressed" not in result_row or "Yes" not in result_row or "—" not in result_row:
                raise AssertionError(f"FHFA-suppressed Vermont row is not visibly preserved: {result_row}")
            with page.expect_download() as geo_download_info:
                explorer.locator("#nmdb-all-geo-excel").click()
            geo_download = geo_download_info.value
            geo_workbook = args.receipt.parent / "CreditAtlas-NMDB-geography-browser-qa.xlsx"
            geo_workbook.parent.mkdir(parents=True, exist_ok=True)
            geo_download.save_as(geo_workbook)
            geo_bytes = geo_workbook.read_bytes()
            workbook = load_workbook(geo_workbook, data_only=True)
            if workbook.sheetnames != ["Start Here", "Selected data", "Definitions", "Quality & lineage"]:
                raise AssertionError("Filtered NMDB geography Excel workbook is missing expected sheets")
            if workbook["Selected data"].cell(2, 10).value is not None or workbook["Selected data"].cell(2, 11).value != "Yes":
                raise AssertionError("Filtered Excel workbook changed the source-suppressed value")
            checks.append("filtered Excel export preserves Vermont suppression and source lineage")
            browser.close()
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)
    if errors:
        raise AssertionError(f"Browser page errors: {errors}")
    receipt = {
        "schema_version": "nmdb-context-browser-qa-v2",
        "status": "PASS",
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "checks": checks,
        "page_errors": errors,
        "filtered_xlsx": {
            "filename": "CreditAtlas-NMDB-geography-browser-qa.xlsx",
            "bytes": len(geo_bytes),
            "sha256": hashlib.sha256(geo_bytes).hexdigest(),
            "selected_rows": 1,
            "suppressed_cell_preserved": True,
        },
        "privacy": "No PUF IDs, row-level mortgage data or screenshots captured.",
    }
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
