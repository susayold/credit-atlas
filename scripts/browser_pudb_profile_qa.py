"""Browser and workbook acceptance for the private FHFA PUDB analyst profile."""

from __future__ import annotations

import argparse
import json
import threading
from datetime import datetime, timezone
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from credit_atlas.server import handler_for
from credit_atlas.service import ResearchService


def _check_workbook(path: Path, profile: dict) -> dict:
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        expected_sheets = ["Overview", "Category Mix", "Missingness", "Data Dictionary", "Sources & Methods"]
        if values.sheetnames != expected_sheets:
            raise AssertionError(f"Unexpected PUDB workbook sheets: {values.sheetnames}")
        overview = values["Overview"]
        if overview["B5"].value != profile["source_release_date"] or overview["B8"].value != profile["rows"]:
            raise AssertionError("PUDB workbook overview does not match the source profile")
        if overview["D5"].value != "Fannie Mae" or overview["E5"].value != profile["enterprise_counts"]["1"]:
            raise AssertionError("PUDB workbook Fannie Mae total does not reconcile")
        if overview["D6"].value != "Freddie Mac" or overview["E6"].value != profile["enterprise_counts"]["2"]:
            raise AssertionError("PUDB workbook Freddie Mac total does not reconcile")

        category_rows = list(values["Category Mix"].iter_rows(min_row=6, values_only=True))
        category_rows = [row for row in category_rows if row and row[0] in profile["category_counts"]]
        if len(category_rows) != sum(len(items) for items in profile["category_counts"].values()):
            raise AssertionError("PUDB category mix workbook row count does not reconcile")
        suppressed_categories = 0
        for row in category_rows:
            field = row[0]
            code = str(row[2]) if row[2] is not None else "<BLANK>"
            count = profile["category_counts"][field][code]
            if count < 30:
                suppressed_categories += 1
                if row[4] != "Suppressed (<30)" or row[5] is not None or row[7] is not None:
                    raise AssertionError(f"Small PUDB category {field}/{code} was not safely suppressed")
                if row[8] != "Suppressed for small count":
                    raise AssertionError(f"Suppression status missing for {field}/{code}")
            elif row[4] != count:
                raise AssertionError(f"PUDB category count differs for {field}/{code}")
        if suppressed_categories != sum(count < 30 for items in profile["category_counts"].values() for count in items.values()):
            raise AssertionError("PUDB category suppression total is inconsistent")

        missing_rows = list(values["Missingness"].iter_rows(min_row=6, values_only=True))
        missing_rows = [row for row in missing_rows if row and row[0] in profile["blank_counts"]]
        if len(missing_rows) != profile["columns"]:
            raise AssertionError("PUDB missingness workbook does not contain all 66 source fields")
        for row in missing_rows:
            expected = profile["blank_counts"][row[0]]
            displayed = row[3]
            if expected < 30 and expected > 0:
                if displayed != "Suppressed (<30)" or row[4] is not None or row[5] is not None:
                    raise AssertionError(f"Small positive missingness count was not suppressed for {row[0]}")
            elif displayed != expected:
                raise AssertionError(f"PUDB blank count differs for {row[0]}")

        dictionary_rows = list(values["Data Dictionary"].iter_rows(min_row=5, values_only=True))
        dictionary_rows = [row for row in dictionary_rows if row and row[1] is not None]
        if len(dictionary_rows) != profile["columns"]:
            raise AssertionError("PUDB dictionary workbook does not contain one row per source field")
        if formulas["Category Mix"]["F6"].value is None or not str(formulas["Category Mix"]["F6"].value).startswith("="):
            raise AssertionError("PUDB category share cells are not editable Excel formulas")

        return {
            "sheet_count": len(expected_sheets),
            "category_rows": len(category_rows),
            "category_suppressed_cells": suppressed_categories,
            "missingness_fields": len(missing_rows),
            "dictionary_fields": len(dictionary_rows),
            "excel_formula_present": True,
        }
    finally:
        values.close()
        formulas.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, type=Path, help="Private workspace holding verified PUDB inputs and workbook")
    parser.add_argument("--web-root", type=Path, help="Public checkout UI directory when source code and data are separate")
    parser.add_argument("--receipt", required=True, type=Path, help="Private QA JSON output under artifacts/private")
    args = parser.parse_args()
    root = args.root.resolve()
    output = args.receipt.resolve().parent
    private_root = (root / "artifacts" / "private").resolve()
    if not output.is_relative_to(private_root):
        raise ValueError("PUDB browser QA receipt and downloaded files must stay under artifacts/private")
    output.mkdir(parents=True, exist_ok=True)

    service = ResearchService(root, review_path=output / "qa-reviews.sqlite3")
    profile = service.fhfa_pudb_profile()
    if profile.get("status") != "AVAILABLE" or profile.get("public_release_approved") is not False:
        raise AssertionError("PUDB profile must be locally available and remain fenced from public release")
    runs = [row for row in service.runs()["runs"] if row.get("descriptive_status") == "verified_rebuild"]
    if not runs:
        raise AssertionError("A verified D01 run is required to reach the mortgage sources workspace")
    case_id = runs[0]["case_id"]

    web_root = args.web_root.resolve() if args.web_root else None
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler_for(service, web_root=web_root))
    server.daemon_threads = True
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    checks: list[str] = []
    errors: list[str] = []
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="msedge", headless=True)
            page = browser.new_page(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            page.set_default_timeout(30_000)
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.on("console", lambda message: errors.append(message.text) if message.type == "error" else None)
            response = page.goto(
                f"http://127.0.0.1:{server.server_port}/?run={case_id}&view=sources",
                wait_until="networkidle",
                timeout=120_000,
            )
            if response is None or response.status != 200:
                raise RuntimeError("Mortgage source page did not return HTTP 200")
            box = page.locator("#fhfa-pudb-acquisition")
            expect(box).to_be_visible(timeout=60_000)
            for value in (f"{profile['rows']:,}", str(profile["columns"]), f"{profile['unique_record_keys']:,}", profile["archive_sha256"]):
                expect(box).to_contain_text(value)
            checks.append("PUDB browser profile shows the receipt-bound release, row/field/key totals and archive hash")

            small_category_count = sum(count < 30 for items in profile["category_counts"].values() for count in items.values())
            details = box.locator("details")
            details.locator("summary").click()
            suppressed_rows = details.get_by_text("Suppressed (<30)", exact=True)
            expect(suppressed_rows).to_have_count(small_category_count)
            enterprise = details.locator("section.card").filter(has=page.get_by_role("heading", name="Enterprise", exact=True))
            expect(enterprise).to_contain_text("987,450")
            expect(enterprise).to_contain_text("1,106,317")
            federal = details.locator("section.card").filter(has=page.get_by_role("heading", name="Federal guarantee", exact=True))
            va_row = federal.locator("tbody tr").filter(has_text="VA-guaranteed")
            expect(va_row).to_contain_text("Suppressed (<30)")
            if "17" in va_row.inner_text():
                raise AssertionError("Suppressed VA count leaked into the browser view")
            checks.append("Category tables reconcile to profile; all sub-30 category counts and shares are suppressed")

            api_response = page.request.get(f"http://127.0.0.1:{server.server_port}/api/fhfa-pudb-profile")
            if api_response.status != 200 or api_response.json() != profile:
                raise AssertionError("PUDB profile API differs from the integrity-checked service output")
            with page.expect_download() as json_download:
                box.get_by_role("button", name="Xuất FHFA PUDB aggregate profile", exact=True).click()
            json_file = output / "FHFA-Enterprise-PUDB-2025-private-acquisition-profile.json"
            json_download.value.save_as(json_file)
            if json.loads(json_file.read_text(encoding="utf-8")) != profile:
                raise AssertionError("PUDB JSON export does not match the verified profile")
            checks.append("Profile API and JSON download match; release fence remains private/not-public-approved")

            workbook_link = box.get_by_role("link", name="Tải workbook Excel phân tích PUDB", exact=True)
            with page.expect_download() as workbook_download:
                workbook_link.click()
            workbook_file = output / "CreditAtlas-FHFA-PUDB-2025-Analyst-Review.xlsx"
            workbook_download.value.save_as(workbook_file)
            workbook_checks = _check_workbook(workbook_file, profile)
            checks.append("Five-sheet analyst workbook reconciles 126 categories and all 66 fields; sub-30 counts suppressed")

            page.set_viewport_size({"width": 390, "height": 844})
            if not page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"):
                raise AssertionError("Mortgage sources / PUDB profile page overflows at 390px")
            checks.append("PUDB profile remains usable at a 390px viewport")
            browser.close()
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)
    if errors:
        raise AssertionError(f"PUDB browser page or console errors: {errors}")

    receipt = {
        "schema_version": "fhfa-pudb-browser-qa-v1",
        "status": "PASS",
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "case_id": case_id,
        "checks": checks,
        "page_errors": errors,
        "data": {
            "source_release_date": profile["source_release_date"],
            "rows": profile["rows"],
            "columns": profile["columns"],
            "unique_record_keys": profile["unique_record_keys"],
            "archive_sha256": profile["archive_sha256"],
            **workbook_checks,
            "public_release_approved": profile["public_release_approved"],
        },
        "limitations": [
            "Workbook was checked as an XLSX package; Microsoft Excel desktop open/recalculate is not asserted.",
            "PUDB is an annual Enterprise acquisition cross-section, not mortgage payment performance or a default target.",
            "Receipt, private JSON profile, QA workbook and screenshots are not copied into the public repository.",
        ],
    }
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
