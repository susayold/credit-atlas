"""Curate Treasury HAMP quarterly cohort tables without imputing blank cells."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pypdf import PdfReader


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_INDEX = REPO_ROOT / "docs/data/treasury-hamp-quarterly-summary-sources-2026-10-17.json"
CURATED_CSV_NAME = "treasury-hamp-quarterly-performance-2018-2023.csv"
CURATED_RECEIPT_NAME = "treasury-hamp-quarterly-performance-receipt.json"
PROGRAMS = [
    ("HAMP Tier 1", "performance of hamp permanent modifications - tier 1", list(range(2009, 2018))),
    ("HAMP Tier 2", "performance of hamp permanent modifications - tier 2", list(range(2012, 2018))),
    ("Streamline HAMP", "performance of hamp permanent modifications - streamline hamp", [2016, 2017]),
]
HORIZONS = [6, 12, 24, 36, 48, 60, 72]
ROW_RE = re.compile(r"^\s*(\d{4}|All)\s+([\d,]+)\b(.*)$")
PCT_RE = re.compile(r"(\d+(?:\.\d+)?)%")
CSV_FIELDS = [
    "source_id", "publication_quarter", "updated_through_date", "program_segment",
    "modification_effective_year", "permanent_modifications_started_as_reported",
    "modification_age_months", "delinquency_60_plus_pct", "delinquency_90_plus_pct",
    "cell_status", "source_pdf", "source_url", "source_pdf_sha256",
]


class CurationError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_text(value: str) -> str:
    return " ".join(value.replace("–", "-").replace("—", "-").split()).casefold()


def parse_summary_text(layout_text: str, standard_text: str, report: dict, source_sha256: str) -> list[dict]:
    """Parse a one-page Treasury cohort summary; trailing absent cells stay null."""
    expected_date = datetime.strptime(report["updated_through"], "%Y-%m-%d").date()
    date_match = re.search(r"Updated through\s+([A-Za-z]+\s+\d{1,2},\s+\d{4})", standard_text, re.IGNORECASE)
    if not date_match:
        raise CurationError(f"Missing Treasury data cutoff in {report['source_id']}")
    observed_date = datetime.strptime(date_match.group(1), "%B %d, %Y").date()
    if observed_date != expected_date:
        raise CurationError(f"Source-index cutoff disagrees with PDF for {report['source_id']}")

    program_rows: dict[str, list[tuple[str, int, list[float]]]] = {}
    current_program: str | None = None
    for line in layout_text.splitlines():
        normalized = normalize_text(line)
        for program, marker, _ in PROGRAMS:
            if marker in normalized:
                current_program = program
                program_rows.setdefault(program, [])
                break
        if current_program is None:
            continue
        match = ROW_RE.match(line)
        if not match:
            continue
        cohort, count, tail = match.groups()
        percentages = [float(value) for value in PCT_RE.findall(tail)]
        if not percentages or len(percentages) > 14 or len(percentages) % 2:
            raise CurationError(f"Malformed percentage row in {report['source_id']}: {line.strip()}")
        program_rows[current_program].append((cohort, int(count.replace(",", "")), percentages))

    output: list[dict] = []
    for program, _, expected_years in PROGRAMS:
        rows = program_rows.get(program, [])
        cohort_keys = [key for key, _, _ in rows]
        observed_years = sorted(int(key) for key in cohort_keys if key != "All")
        if observed_years != expected_years or len(cohort_keys) != len(set(cohort_keys)):
            raise CurationError(f"Cohort years or uniqueness differ for {report['source_id']} / {program}")
        for cohort, modifications, percentages in rows:
            if modifications <= 0:
                raise CurationError(f"Non-positive modification count in {report['source_id']}")
            for age_index, age in enumerate(HORIZONS):
                offset = age_index * 2
                p60 = percentages[offset] if offset < len(percentages) else None
                p90 = percentages[offset + 1] if offset + 1 < len(percentages) else None
                if (p60 is None) != (p90 is None):
                    raise CurationError(f"Only one delinquency metric is present at a horizon in {report['source_id']}")
                if p60 is not None and not (0 <= p90 <= p60 <= 100):
                    raise CurationError(f"Delinquency rates outside valid subset/range in {report['source_id']}")
                output.append({
                    "source_id": report["source_id"],
                    "publication_quarter": report["quarter"],
                    "updated_through_date": report["updated_through"],
                    "program_segment": program,
                    "modification_effective_year": cohort,
                    "permanent_modifications_started_as_reported": modifications,
                    "modification_age_months": age,
                    "delinquency_60_plus_pct": p60,
                    "delinquency_90_plus_pct": p90,
                    "cell_status": "reported_in_source" if p60 is not None else "not_available_in_published_table",
                    "source_pdf": report["file"],
                    "source_url": report["url"],
                    "source_pdf_sha256": source_sha256,
                })
    if not output:
        raise CurationError(f"No HAMP cohort rows parsed for {report['source_id']}")
    return output


def parse_pdf(path: Path, report: dict) -> tuple[list[dict], str, int]:
    if not path.is_file():
        raise CurationError(f"Missing private Treasury PDF: {path}")
    digest = sha256_file(path)
    reader = PdfReader(path, strict=True)
    if len(reader.pages) != 1:
        raise CurationError(f"Expected one-page Treasury summary: {path.name}")
    page = reader.pages[0]
    layout_text = page.extract_text(extraction_mode="layout") or ""
    standard_text = page.extract_text() or ""
    rows = parse_summary_text(layout_text, standard_text, report, digest)
    return rows, digest, len(reader.pages)


def style_table(sheet, name: str, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    ref = sheet.dimensions
    if ref and ":" in ref:
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                              showColumnStripes=False, showFirstColumn=False,
                                              showLastColumn=False)
        sheet.add_table(table)
    for column_cells in sheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max((len(str(cell.value or "")) for cell in column_cells[:250]), default=10)
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 46)


def write_workbook(rows: list[dict], reports: list[dict], target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cohort KPIs"
    ws.append(CSV_FIELDS)
    for row in rows:
        ws.append([row[field] for field in CSV_FIELDS])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (8, 9) and cell.value is not None:
                cell.number_format = "0.0\"%\""
    style_table(ws, "HampQuarterlyCohortKpis")

    source_ws = wb.create_sheet("Source Index")
    source_fields = ["source_id", "publication_quarter", "published_on", "updated_through", "file", "url", "bytes", "sha256", "page_count"]
    source_ws.append(source_fields)
    for report in reports:
        source_ws.append([report[field] for field in source_fields])
    style_table(source_ws, "HampQuarterlySources")
    for row in source_ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].hyperlink = row[0].value
        row[0].style = "Hyperlink"

    guide = wb.create_sheet("Read Me")
    guide.append(["Topic", "Definition / limitation"])
    notes = [
        ("Source", "Official U.S. Treasury HAMP performance summaries linked from the Making Home Affordable Program Performance Report page."),
        ("Unit of observation", "Published snapshot × HAMP program segment × modification effective-year cohort × loan-age milestone."),
        ("Rates", "Treasury-reported 60+ days delinquent and 90+ days delinquent percentages at month 6, 12, 24, 36, 48, 60 and 72 after modification."),
        ("90+ caveat", "Treasury states payment status is not reported after program disqualification at 90+ days delinquent; 90+ loans remain in future 90+ metrics even if they later cure or pay off."),
        ("Denominator caveat", "Paid-off or withdrawn loans leave future periods, so the reported permanent-modification count can change between snapshots. The denominator is not a fixed original cohort."),
        ("Blank cells", "No number in the published table is preserved as blank/status not_available_in_published_table. It is not zero and is not imputed."),
        ("Missing report", "Treasury's current index lists no 2021-Q3 HAMP performance summary. No value is inserted for that quarter."),
        ("Use", "Descriptive program-specific cohort surveillance only. These measures are not a general mortgage PD/default label and do not qualify a prediction model."),
        ("Relationship to monthly file", "These aggregate cohort snapshots do not reconstruct the missing 2018-2021 monthly loan activity in the separate HAMP supplemental ZIP."),
        ("Forecasting", "No forecasts, fills, or projections are made in this workbook."),
    ]
    for row in notes:
        guide.append(row)
    style_table(guide, "HampQuarterlyReadme")
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 110
    for row in guide.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row[0].row].height = 42
    wb.save(target)
    check = load_workbook(target, read_only=True, data_only=True)
    if check["Cohort KPIs"].max_row != len(rows) + 1 or check["Source Index"].max_row != len(reports) + 1:
        raise CurationError("Written workbook row counts fail round-trip validation")
    check.close()


def curate(root: Path, output_dir: Path, workbook_path: Path) -> dict:
    root = root.resolve()
    if root.drive.upper() != "D:":
        raise CurationError("Treasury PDFs, workbook and receipts must stay on D:")
    expected_private_root = (root / "data/private/mortgage/treasury-hamp/quarterly-summary").resolve()
    output_dir = output_dir.resolve()
    workbook_path = workbook_path.resolve()
    if not output_dir.is_relative_to(expected_private_root) or not workbook_path.is_relative_to(expected_private_root):
        raise CurationError("Curated values and workbook must remain in the private Treasury HAMP folder on D:")
    index = json.loads(SOURCE_INDEX.read_text(encoding="utf-8"))
    if index.get("source_count") != len(index["reports"]):
        raise CurationError("Quarterly source index count does not reconcile")
    pdf_dir = root / "data/private/mortgage/treasury-hamp/source_docs/quarterly_summaries"
    all_rows: list[dict] = []
    source_profiles = []
    for report in index["reports"]:
        rows, digest, page_count = parse_pdf(pdf_dir / report["file"], report)
        all_rows.extend(rows)
        path = pdf_dir / report["file"]
        source_profiles.append({
            "source_id": report["source_id"], "publication_quarter": report["quarter"],
            "published_on": report["published_on"], "updated_through": report["updated_through"], "file": report["file"],
            "url": report["url"], "bytes": path.stat().st_size, "sha256": digest,
            "page_count": page_count, "parsed_long_rows": len(rows),
        })
    unique_keys = {(r["source_id"], r["program_segment"], r["modification_effective_year"], r["modification_age_months"])
                   for r in all_rows}
    if len(unique_keys) != len(all_rows):
        raise CurationError("Output contains duplicate report/program/cohort/horizon keys")

    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / CURATED_CSV_NAME
    partial = csv_path.with_suffix(csv_path.suffix + ".partial")
    if partial.exists():
        raise CurationError("Partial curated CSV exists; inspect it before retry")
    try:
        with partial.open("x", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS, extrasaction="raise")
            writer.writeheader()
            writer.writerows(all_rows)
        partial.replace(csv_path)
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    write_workbook(all_rows, source_profiles, workbook_path)

    missing_cells = sum(1 for row in all_rows if row["cell_status"] != "reported_in_source")
    receipt = {
        "schema_version": 1,
        "status": "PASS",
        "curation_id": "treasury-hamp-quarterly-curation-" + uuid.uuid4().hex[:12],
        "completed_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_page": index["source_page"],
        "source_count": len(source_profiles),
        "published_long_rows": len(all_rows),
        "published_rate_pairs": len(all_rows) - missing_cells,
        "blank_rate_pairs_preserved": missing_cells,
        "source_periods": [r["publication_quarter"] for r in source_profiles],
        "unlisted_periods": ["2021-Q3"],
        "csv_file": CURATED_CSV_NAME,
        "csv_bytes": csv_path.stat().st_size,
        "csv_sha256": sha256_file(csv_path),
        "workbook_private_path": str(workbook_path.resolve()),
        "workbook_sha256": sha256_file(workbook_path),
        "method": "Direct transcription of published 60+/90+ delinquency percentages by reported age horizon; no calculation, interpolation, imputation or forecast.",
        "limitations": [
            "Aggregate HAMP program cohorts only; not a general mortgage portfolio or representative borrower sample.",
            "90+ rates carry forward loans disqualified for 90+ delinquency even if later cured or paid off.",
            "Paid-off/withdrawn loans leave later periods, so the published count is not a fixed denominator.",
            "A blank source cell remains unavailable; it is not zero or a negative event.",
            "Quarterly cohort snapshots do not fill the 2018-2021 loan-level monthly activity gap in the separate ZIP.",
            "No PD model was fit or qualified.",
        ],
        "sources": source_profiles,
        "public_release_approved": False,
        "visibility": "aggregate_public_source_curated_data",
    }
    receipt_path = output_dir / CURATED_RECEIPT_NAME
    receipt_partial = receipt_path.with_suffix(receipt_path.suffix + ".partial")
    if receipt_partial.exists():
        raise CurationError("Partial receipt exists; inspect it before retry")
    try:
        with receipt_partial.open("x", encoding="utf-8") as stream:
            json.dump(receipt, stream, indent=2, ensure_ascii=False, allow_nan=False)
            stream.write("\n")
        receipt_partial.replace(receipt_path)
    except Exception:
        receipt_partial.unlink(missing_ok=True)
        raise
    return {"status": receipt["status"], "source_count": len(source_profiles),
            "published_long_rows": len(all_rows), "published_rate_pairs": len(all_rows) - missing_cells,
            "blank_rate_pairs_preserved": missing_cells, "csv_sha256": receipt["csv_sha256"],
            "workbook_sha256": receipt["workbook_sha256"]}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(r"D:\data\Codex\credit-atlas"))
    parser.add_argument("--output-dir", type=Path, default=Path(
        r"D:\data\Codex\credit-atlas\data\private\mortgage\treasury-hamp\quarterly-summary"))
    parser.add_argument("--workbook", type=Path, default=Path(
        r"D:\data\Codex\credit-atlas\data\private\mortgage\treasury-hamp\quarterly-summary\CreditAtlas-Treasury-HAMP-Quarterly-Summary-2018-2023.xlsx"))
    args = parser.parse_args()
    try:
        result = curate(args.root, args.output_dir, args.workbook)
        print(json.dumps(result, indent=2, ensure_ascii=False))
    except (OSError, CurationError, ValueError) as exc:
        parser.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
