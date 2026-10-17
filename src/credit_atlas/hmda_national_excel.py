"""Query-scoped Excel export for receipt-verified 2025 HMDA aggregates."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "EEF3F8"
LOAN_TYPES = {
    "1": "Conventional",
    "2": "FHA",
    "3": "VA",
    "4": "RHS/FSA",
}


def workbook_bytes(service, query: dict) -> bytes:
    """Build an analyst workbook from the selected state/product aggregate only."""
    report = service.hmda_national_views.report(query.get("view_id", ""), query)
    filters = report.get("applied_filters")
    summary = report.get("filtered_summary")
    if not filters or not summary:
        raise ValueError("Select a state or loan type before exporting the filtered HMDA workbook")
    actions = summary.get("actions", [])
    if (len(actions) != 8 or {row["code"] for row in actions} != {str(i) for i in range(1, 9)}
            or sum(row["records"] for row in actions) != summary["records"]):
        raise ValueError("Filtered HMDA action profile is incomplete or does not reconcile")

    book = Workbook()
    start = book.active
    start.title = "Summary"
    start.sheet_view.showGridLines = False
    start.merge_cells("A1:C1")
    start["A1"] = "CreditAtlas · HMDA filtered application summary"
    start["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    start["A1"].fill = PatternFill("solid", fgColor=NAVY)
    start["A1"].alignment = Alignment(vertical="center")
    start.row_dimensions[1].height = 34

    selected_state = filters.get("state_code")
    selected_loan_type = filters.get("loan_type")
    state_label = "All states / territories" if selected_state is None else (
        "Source blank" if selected_state == "(SOURCE_BLANK)" else selected_state)
    product_label = "All loan types" if selected_loan_type is None else (
        f"{LOAN_TYPES.get(selected_loan_type, 'Other / source code')} · {selected_loan_type}")
    metadata = [
        ("Activity year", report["activity_year"]),
        ("Release", report["release_type"]),
        ("Snapshot / as-of date", report["snapshot_date"]),
        ("Selected state / territory", state_label),
        ("Selected loan type", product_label),
        ("Publisher", "FFIEC / Consumer Financial Protection Bureau"),
    ]
    for row_index, (label, value) in enumerate(metadata, start=3):
        start.cell(row_index, 1, label)
        start.cell(row_index, 2, value)
    start["A10"] = "Key performance metrics"
    start["A10"].font = Font(bold=True, color="FFFFFF")
    start["A10"].fill = PatternFill("solid", fgColor=NAVY)
    for cell in start[10][:3]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for col, label in enumerate(("Metric", "Value", "Definition"), start=1):
        cell = start.cell(11, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
    metrics = [
        ("LAR records · actions 1–8", summary["records"], "All in-scope public LAR rows in the selected state/product filter."),
        ("Decision denominator · actions 1+2+3", summary["decision_denominator"], "Originated + approved, not accepted + denied applications."),
        ("Denied numerator · action 3", summary["denied_numerator"], "Applications reported as denied within this filtered scope."),
        ("Descriptive denied share", (float(Decimal(summary["descriptive_denial_rate"]))
            if summary["descriptive_denial_rate"] is not None else "Unavailable"),
            "Action 3 divided by actions 1+2+3; not repayment default or PD."),
        ("Public loan amount midpoint sum · USD", float(Decimal(summary["public_amount_sum_usd"])), "Sum of public interval midpoints; not funded cash or exact origination dollars."),
    ]
    for row_index, metric in enumerate(metrics, start=12):
        for col_index, value in enumerate(metric, start=1):
            start.cell(row_index, col_index, value)
    start["B15"].number_format = "0.0000%"
    start["B16"].number_format = '"$"#,##0.00;[Red]("$"#,##0.00);–'

    actions_sheet = book.create_sheet("Actions")
    actions_sheet.sheet_view.showGridLines = False
    headers = ["Action code", "HMDA disposition", "Records", "Share of filtered LAR", "In decision denominator"]
    actions_sheet.append(headers)
    for row in actions:
        actions_sheet.append([
            row["code"], row["name"], row["records"], None,
            "Yes · action 1–3" if row["code"] in {"1", "2", "3"} else "No",
        ])
    for row_index in range(2, 10):
        actions_sheet.cell(row_index, 4, f"=C{row_index}/SUM($C$2:$C$9)")
        actions_sheet.cell(row_index, 4).number_format = "0.0000%"
    actions_table = Table(displayName="HMDAFilteredActions", ref="A1:E9")
    actions_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    actions_sheet.add_table(actions_table)
    actions_sheet.freeze_panes = "A2"

    checks = [
        ("All action counts sum to filtered LAR records", '=IF(SUM(Actions!$C$2:$C$9)=B12,"PASS","CHECK")'),
        ("Actions 1–3 sum to decision denominator", '=IF(SUM(Actions!$C$2:$C$4)=B13,"PASS","CHECK")'),
        ("Action 3 equals denied numerator", '=IF(Actions!$C$4=B14,"PASS","CHECK")'),
        ("Denied share matches numerator / denominator", '=IF(B13=0,"N/A",IF(ABS(B15-B14/B13)<1E-12,"PASS","CHECK"))'),
    ]
    start["A18"] = "Reconciliation checks"
    start["A18"].font = Font(bold=True, color="FFFFFF")
    start["A18"].fill = PatternFill("solid", fgColor=NAVY)
    for row_index, (label, formula) in enumerate(checks, start=19):
        start.cell(row_index, 1, label)
        start.cell(row_index, 2, formula)
        start.cell(row_index, 2).fill = PatternFill("solid", fgColor=PALE)

    method = book.create_sheet("Method & sources")
    method.sheet_view.showGridLines = False
    method.append(["Topic", "Definition / source note"])
    method_rows = [
        ("Selected filters", "; ".join(f"{key}={value}" for key, value in sorted(filters.items()))),
        ("Question answered", "How did reported application dispositions differ inside the selected state and/or loan-type population?"),
        ("Numerator", "HMDA action_taken = 3 (denied)."),
        ("Denominator", "HMDA action_taken in 1, 2, or 3 (originated, approved not accepted, or denied)."),
        ("Interpretation", "Descriptive underwriting/application disposition rate. It is not borrower repayment performance, default, delinquency, PD, or causal impact."),
        ("Amount measure", "HMDA public loan_amount values are published interval midpoints; their sum is not funded cash."),
        ("Rows included", "Public aggregate counts only. No borrower-level rows, LEIs, or raw LAR records are exported."),
        ("Source release", f"{report['release_type']} · activity year {report['activity_year']} · as of {report['snapshot_date']}"),
        ("Official release page", report["source_url"]),
        ("Official API documentation", "https://ffiec.cfpb.gov/documentation/api/data-browser/"),
        ("Visibility", "This local report is generated from a receipt-verified private workspace source. Raw source and detailed receipts remain private; public release is not approved."),
        *[("Use limitation", item) for item in report["limitations"]],
    ]
    for row in method_rows:
        method.append(row)
    method.freeze_panes = "A2"

    for sheet in (start, actions_sheet, method):
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    start.column_dimensions["A"].width = 43
    start.column_dimensions["B"].width = 34
    start.column_dimensions["C"].width = 82
    start.freeze_panes = "A11"
    for row_index in range(2, start.max_row + 1):
        start.row_dimensions[row_index].height = 32
    for col, width in {"A": 16, "B": 38, "C": 20, "D": 26, "E": 28}.items():
        actions_sheet.column_dimensions[col].width = width
    method.column_dimensions["A"].width = 28
    method.column_dimensions["B"].width = 112
    for row_index in range(2, method.max_row + 1):
        method.row_dimensions[row_index].height = 42
    for cell in method[10]:
        if cell.column == 2:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"

    output = BytesIO()
    book.save(output)
    payload = output.getvalue()
    check = load_workbook(BytesIO(payload), data_only=False, read_only=False)
    if (check.sheetnames != ["Summary", "Actions", "Method & sources"]
            or check["Actions"].max_row != 9
            or sum(1 for row in check["Summary"].iter_rows() for cell in row
                   if isinstance(cell.value, str) and cell.value.startswith("=")) != 4
            or sum(1 for row in check["Actions"].iter_rows() for cell in row
                   if isinstance(cell.value, str) and cell.value.startswith("=")) != 8):
        raise ValueError("HMDA filtered workbook failed its structure/formula check")
    check.close()
    return payload
