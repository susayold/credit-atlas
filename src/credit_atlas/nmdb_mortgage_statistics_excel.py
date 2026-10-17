"""Query-scoped Excel export for FHFA NMDB mortgage-market aggregates."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


NAVY = "243B53"
BLUE = "486581"
PALE = "EAF1F7"


def workbook_bytes(service, query: dict) -> bytes:
    report = service.nmdb_mortgage_statistics(query)
    book = Workbook()
    start = book.active
    start.title = "Start Here"
    start.sheet_view.showGridLines = False
    info = [
        ["CreditAtlas · FHFA NMDB market statistics", ""],
        ["Selected product", report["product_label"]],
        ["Product interpretation", report["definition"]],
        ["Selected geography", report["filters"]["geo_level"] + " · " + report["filters"]["geo_name"] + " (" + report["filters"]["geo_id"] + ")"],
        ["Mortgage market", report["filters"]["market"]],
        ["Selected series", report["series"]["series_id"] + " · " + report["series"]["label"]],
        ["Series definition", report["series"]["definition"]],
        ["Series unit", report["series"]["unit"]],
        ["FHFA-supported weights", ", ".join(report["supported_weights"]) + " · " + report["weight_note"]],
        ["Selected weighting", report["measure"] + " · " + report["measure_definition"]],
        ["Coverage", f"{report['filters']['from_year']}–{report['filters']['to_year']} · {report['frequency']}"],
        ["Source release / retrieved", report["source_release_date"] + " · " + report["retrieved_at"]],
        ["Mart SHA-256", report["mart_sha256"]],
        ["Use limitation", "Public aggregate market context. Not loan-level data, a PD estimate, Freddie vintage, approval rate, or model validation sample."],
        ["Missing / suppressed", "Blank, dot, and SUPPRESSED are preserved. Do not interpret unavailable cells as zero; no forecast or imputation is applied."],
        ["Weight interpretation", "VALUE1: count weighted. VALUE2: origination-dollar weighted for New; quarter-end UPB weighted for Outstanding. These are different estimands."],
        ["FHFA source", "https://www.fhfa.gov/data/nmdb"],
        ["FHFA dictionary", "https://www.fhfa.gov/document/d/nmdb/nmdb-aggregate-statistics-data-dictionary-technical-notes"],
    ]
    start.append(info[0])
    for row in info[1:]:
        start.append(row)
    start.merge_cells("A1:B1")
    start["A1"].fill = PatternFill("solid", fgColor=NAVY)
    start["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    start["A1"].alignment = Alignment(vertical="center")
    start.row_dimensions[1].height = 34
    for row in start.iter_rows(min_row=2, max_row=len(info)):
        row[0].font = Font(name="Aptos", bold=True, color=NAVY)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    start.column_dimensions["A"].width = 28
    start.column_dimensions["B"].width = 108
    for row_index in range(2, len(info) + 1):
        start.row_dimensions[row_index].height = 34
    for cell in (start["B17"], start["B18"]):
        cell.hyperlink = cell.value
        cell.style = "Hyperlink"
    start.freeze_panes = "A2"

    data = book.create_sheet("Selected data")
    data.sheet_view.showGridLines = False
    headers = [
        "Period", "Year", "Quarter", "Geography level", "Geography ID", "Geography name",
        "Mortgage market", "Series ID", "Series label", "Series unit", "FHFA-supported weights", "Selected weight", "Selected value",
        "VALUE1 raw · count-weighted", "VALUE2 raw · amount / UPB-weighted", "Suppressed",
        "Source", "Frequency",
    ]
    data.append(headers)
    for row in report["rows"]:
        data.append([
            row["period"], row["year"], row["quarter"], row["geography_level"], row["geography_id"],
            row["geography_name"], row["market"], row["series_id"], report["series"]["label"],
            report["series"]["unit"], ", ".join(report["supported_weights"]), report["measure"], row["selected_value"], row["value1_raw"], row["value2_raw"],
            "Yes · fewer than 3 loans" if row["suppressed"] else "No", row["source"], row["frequency"],
        ])
    table = Table(displayName="NMDBSelectedData", ref=f"A1:R{data.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    data.add_table(table)
    data.freeze_panes = "A2"
    for cell in data[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data.row_dimensions[1].height = 42
    widths = [15, 10, 10, 20, 14, 28, 36, 22, 42, 18, 24, 16, 18, 22, 28, 22, 12, 14]
    for index, width in enumerate(widths, start=1):
        data.column_dimensions[get_column_letter(index)].width = width
    for cell in data["M"][1:]:
        cell.number_format = '0.00"%";[Red](0.00"%");–' if report["series"]["unit"] == "%" else "#,##0.00;[Red](#,##0.00);–"
    for row in data.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if len(report["rows"]) > 1 and any(row["selected_value"] is not None for row in report["rows"]):
        chart = LineChart()
        chart.title = report["series"]["label"] + " (" + report["measure"] + ")"
        chart.y_axis.title = report["series"]["unit"]
        chart.x_axis.title = "Published period"
        chart.style = 13
        chart.height = 8.2
        chart.width = 20
        chart.display_blanks = "gap"
        chart.add_data(Reference(data, min_col=13, min_row=1, max_row=data.max_row), titles_from_data=True)
        chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=data.max_row))
        chart.series[0].cat = AxDataSource(
            strRef=StrRef(f"'Selected data'!$A$2:$A${data.max_row}")
        )
        data.add_chart(chart, "T2")

    codebook = book.create_sheet("Definitions")
    codebook.sheet_view.showGridLines = False
    definitions = [
        ["Field / concept", "Definition"],
        ["SOURCE", "FHFA NMDB aggregate source identifier."],
        ["FREQUENCY / PERIOD", "Annual for New originations; Quarterly for Outstanding stock. PERIOD follows FHFA's year or YYYYQn convention."],
        ["GEOLEVEL / GEOID / GEONAME", "FHFA geography level, source identifier and published name; retain all three for grouping."],
        ["MARKET", "FHFA mortgage market/submarket category. Do not add overlapping parent/submarket rows."],
        ["SERIESID", report["series"]["series_id"] + " · " + report["series"]["label"] + ": " + report["series"]["definition"]],
        ["VALUE1", "Weighted by number of mortgages. Each NMDB sample mortgage represents 20 actual mortgages."],
        ["VALUE2 · New", "Weighted by origination loan amount."],
        ["VALUE2 · Outstanding", "Weighted by unpaid principal balance at quarter end."],
        ["Series-specific support", ", ".join(report["supported_weights"]) + " · " + report["weight_note"]],
        ["SUPPRESSED", "1 means the statistic is based on fewer than three loans. It remains unavailable and is never zero-filled."],
        ["Units", "Use the series-specific source scale shown on Start Here and the FHFA data dictionary. The export does not rescale values."],
        ["Source dictionary", "https://www.fhfa.gov/document/d/nmdb/nmdb-aggregate-statistics-data-dictionary-technical-notes"],
    ]
    for row in definitions:
        codebook.append(row)
    codebook.freeze_panes = "A2"
    codebook.column_dimensions["A"].width = 28
    codebook.column_dimensions["B"].width = 112
    for cell in codebook[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF")
    for row in codebook.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=NAVY)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_index in range(2, codebook.max_row + 1):
        codebook.row_dimensions[row_index].height = 34

    quality = book.create_sheet("Quality & lineage")
    quality.sheet_view.showGridLines = False
    quality_rows = [
        ["Check", "Result"],
        ["Rows in selected extract", report["quality"]["observations"]],
        ["Suppressed rows", report["quality"]["suppressed_rows"]],
        ["Selected value unavailable", report["quality"]["selected_value_unavailable_rows"]],
        ["Forecast or imputation", "None"],
        ["Mart SHA-256", report["mart_sha256"]],
        ["Source release date", report["source_release_date"]],
        ["Retrieved at (UTC)", report["retrieved_at"]],
        ["Source page", "https://www.fhfa.gov/data/nmdb"],
        ["Product separation", "Origination flow and quarter-end stock remain separate products and have not been joined or summed."],
    ]
    for row in quality_rows:
        quality.append(row)
    quality.column_dimensions["A"].width = 32
    quality.column_dimensions["B"].width = 112
    for cell in quality[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF")
    for row in quality.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=NAVY)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_index in range(2, quality.max_row + 1):
        quality.row_dimensions[row_index].height = 32
    for sheet in book.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    book.save(output)
    return output.getvalue()
