"""Read the privately curated Treasury HAMP quarterly cohort summaries."""
from __future__ import annotations

import csv
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


CSV_RELATIVE_PATH = Path(
    "data/private/mortgage/treasury-hamp/quarterly-summary/"
    "treasury-hamp-quarterly-performance-2018-2023.csv"
)
FIELDS = (
    "source_id", "publication_quarter", "updated_through_date", "program_segment",
    "modification_effective_year", "permanent_modifications_started_as_reported",
    "modification_age_months", "delinquency_60_plus_pct", "delinquency_90_plus_pct",
    "cell_status", "source_url", "source_pdf_sha256",
)
RATE_FIELDS = {"delinquency_60_plus_pct", "delinquency_90_plus_pct"}
LIMITATIONS = [
    "Treasury cohort summaries are descriptive program surveillance, not loan-level data or a general mortgage PD/default label.",
    "Published 60+ and 90+ delinquency percentages use a changing reported population; paid-off or withdrawn modifications leave later periods.",
    "Treasury notes that 90+ delinquency is carried into future 90+ metrics after program disqualification, even if the loan later cures or pays off.",
    "Blank published cells remain unavailable. They are not zero, forecast, or imputed.",
    "The 2021-Q3 summary is absent from the official source index. No quarter is inserted for it.",
    "Do not pool rates across modification vintages or program segments; the workbook and app retain each published cohort row.",
    "These aggregate quarterly summaries do not fill the monthly loan-level HAMP activity gap before 2022-01.",
]


def _read_rows(root: Path) -> tuple[list[dict], dict]:
    root = root.resolve()
    if root.drive.upper() != "D:":
        raise ValueError("CreditAtlas private storage must remain on D:")
    path = (root / CSV_RELATIVE_PATH).resolve()
    private_root = (root / "data/private/mortgage/treasury-hamp/quarterly-summary").resolve()
    if not path.is_relative_to(private_root):
        raise ValueError("Treasury HAMP curated data escaped the private storage root")
    if not path.is_file():
        return [], {"available": False, "reason": "Private Treasury HAMP curated data is not installed in this workspace."}
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("Treasury HAMP curated file exceeds the 20 MB serving bound")
    rows = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if tuple(reader.fieldnames or ()) != tuple(FIELDS[:10]) + ("source_pdf", "source_url", "source_pdf_sha256"):
            raise ValueError("Treasury HAMP curated CSV has an unexpected data contract")
        for raw in reader:
            if len(rows) >= 20_000:
                raise ValueError("Treasury HAMP curated row bound exceeded")
            quarter = raw["publication_quarter"]
            segment = raw["program_segment"]
            vintage = raw["modification_effective_year"]
            age = int(raw["modification_age_months"])
            started = int(raw["permanent_modifications_started_as_reported"])
            if (len(quarter) != 7 or quarter[4:6] != "-Q" or quarter[:4].isdigit() is False
                    or segment not in {"HAMP Tier 1", "HAMP Tier 2", "Streamline HAMP"}
                    or (vintage != "All" and (len(vintage) != 4 or not vintage.isdigit()))
                    or age not in {6, 12, 24, 36, 48, 60, 72} or started <= 0):
                raise ValueError("Treasury HAMP curated CSV contains an invalid cohort key")
            values: dict[str, object] = {
                "source_id": raw["source_id"],
                "publication_quarter": quarter,
                "updated_through_date": raw["updated_through_date"],
                "program_segment": segment,
                "modification_effective_year": vintage,
                "permanent_modifications_started_as_reported": started,
                "modification_age_months": age,
                "cell_status": raw["cell_status"],
                "source_url": raw["source_url"],
                "source_pdf_sha256": raw["source_pdf_sha256"],
            }
            if (values["cell_status"] not in {"reported_in_source", "not_available_in_published_table"}
                    or not str(values["source_url"]).startswith("https://home.treasury.gov/")):
                raise ValueError("Treasury HAMP curated CSV contains an invalid source or cell status")
            for field in RATE_FIELDS:
                text = raw[field].strip()
                values[field] = None if not text else float(text)
                if values[field] is not None and not (0 <= values[field] <= 100):
                    raise ValueError("Treasury HAMP curated CSV contains a rate outside 0–100 percent")
            if (values["cell_status"] == "reported_in_source") != (values["delinquency_60_plus_pct"] is not None
                    and values["delinquency_90_plus_pct"] is not None):
                raise ValueError("Treasury HAMP curated CSV status disagrees with the rate cells")
            if values["delinquency_60_plus_pct"] is not None and values["delinquency_90_plus_pct"] > values["delinquency_60_plus_pct"]:
                raise ValueError("Treasury HAMP 90+ rate exceeds the 60+ rate")
            rows.append(values)
    if not rows:
        raise ValueError("Treasury HAMP curated CSV is empty")
    rows.sort(key=lambda row: (
        row["publication_quarter"], row["program_segment"],
        (1, 0) if row["modification_effective_year"] == "All" else (0, int(row["modification_effective_year"])),
        row["modification_age_months"],
    ))
    quarters = sorted({str(row["publication_quarter"]) for row in rows})
    segments = sorted({str(row["program_segment"]) for row in rows})
    vintages = sorted({str(row["modification_effective_year"]) for row in rows if row["modification_effective_year"] != "All"}, key=int)
    if any(row["modification_effective_year"] == "All" for row in rows):
        vintages.append("All")
    ages = sorted({int(row["modification_age_months"]) for row in rows})
    reported = sum(row["cell_status"] == "reported_in_source" for row in rows)
    meta = {
        "available": True,
        "coverage": {
            "first_publication_quarter": quarters[0], "last_publication_quarter": quarters[-1],
            "publication_quarters": len(quarters), "source_reports": len({row["source_id"] for row in rows}),
            "rows": len(rows), "reported_cells": reported, "unavailable_cells": len(rows) - reported,
            "program_segments": segments,
        },
        "filter_options": {
            "publication_quarters": ["all", *quarters],
            "program_segments": ["all", *segments],
            "modification_effective_years": ["all", *vintages],
            "modification_age_months": ["all", *(str(age) for age in ages)],
        },
    }
    return rows, meta


def report(root: Path, query: dict | None = None) -> dict:
    rows, meta = _read_rows(root)
    if not meta["available"]:
        return {**meta, "filters": {}, "rows": [], "quality": {}, "limitations": LIMITATIONS}
    query = query or {}
    allowed = {"publication_quarter", "program_segment", "modification_effective_year", "modification_age_months"}
    if set(query) - allowed:
        raise ValueError("Unknown Treasury HAMP filter")
    filters = {key: str(query.get(key) or "all") for key in sorted(allowed)}
    choices = {
        "publication_quarter": set(meta["filter_options"]["publication_quarters"]),
        "program_segment": set(meta["filter_options"]["program_segments"]),
        "modification_effective_year": set(meta["filter_options"]["modification_effective_years"]),
        "modification_age_months": set(meta["filter_options"]["modification_age_months"]),
    }
    if any(value not in choices[key] for key, value in filters.items()):
        raise ValueError("Treasury HAMP filter value is unavailable")
    selected = [row for row in rows if all(
        filters[key] == "all" or str(row[row_key]) == filters[key]
        for key, row_key in (
            ("publication_quarter", "publication_quarter"),
            ("program_segment", "program_segment"),
            ("modification_effective_year", "modification_effective_year"),
            ("modification_age_months", "modification_age_months"),
        )
    )]
    distinct_cohorts = {(row["publication_quarter"], row["program_segment"], row["modification_effective_year"]) for row in selected}
    source_ids = {row["source_id"] for row in selected}
    return {
        **meta,
        "filters": filters,
        "quality": {
            "selected_cells": len(selected),
            "reported_cells": sum(row["cell_status"] == "reported_in_source" for row in selected),
            "unavailable_cells": sum(row["cell_status"] != "reported_in_source" for row in selected),
            "distinct_cohorts_across_snapshots": len(distinct_cohorts),
            "source_reports": len(source_ids),
        },
        "rows": selected,
        "limitations": LIMITATIONS,
    }


def workbook_bytes(result: dict) -> bytes:
    if not result.get("available"):
        raise FileNotFoundError("Private Treasury HAMP data is unavailable")
    wb = Workbook()
    data = wb.active
    data.title = "Cohort observations"
    fields = [
        ("publication_quarter", "Publication quarter"), ("updated_through_date", "Updated through"),
        ("program_segment", "Program segment"), ("modification_effective_year", "Modification vintage"),
        ("permanent_modifications_started_as_reported", "Modifications started · source"),
        ("modification_age_months", "Modification age · months"),
        ("delinquency_60_plus_pct", "60+ delinquency · %"),
        ("delinquency_90_plus_pct", "90+ delinquency · %"),
        ("cell_status", "Published cell status"), ("source_id", "Source ID"),
        ("source_url", "Treasury source URL"), ("source_pdf_sha256", "Source PDF SHA-256"),
    ]
    data.append([label for _, label in fields])
    for row in result["rows"]:
        data.append([row[key] for key, _ in fields])
    _style(data, "HampQuarterlyFiltered")
    for row in data.iter_rows(min_row=2):
        for cell in (row[6], row[7]):
            if cell.value is not None:
                cell.number_format = '0.0"%"'
        row[10].hyperlink = row[10].value
        row[10].style = "Hyperlink"

    sources = wb.create_sheet("Source index")
    sources.append(["Source ID", "Quarter", "Updated through", "Official URL"])
    unique = {row["source_id"]: row for row in result["rows"]}
    for source_id, row in sorted(unique.items()):
        sources.append([source_id, row["publication_quarter"], row["updated_through_date"], row["source_url"]])
    _style(sources, "HampQuarterlySelectedSources")
    for row in sources.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].hyperlink = row[0].value
        row[0].style = "Hyperlink"

    guide = wb.create_sheet("Method & limits")
    guide.append(["Topic", "Definition / interpretation"])
    guide.append(["Selected filters", "; ".join(f"{key}={value}" for key, value in result["filters"].items())])
    guide.append(["Selected rows", result["quality"]["selected_cells"]])
    guide.append(["Reported / unavailable cells", f'{result["quality"]["reported_cells"]} / {result["quality"]["unavailable_cells"]}'])
    for limitation in result["limitations"]:
        guide.append(["Use limitation", limitation])
    _style(guide, "HampQuarterlyMethod")
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 110
    for row in guide.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row[0].row].height = 38
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    check = load_workbook(output, read_only=True, data_only=True)
    if check["Cohort observations"].max_row != len(result["rows"]) + 1:
        raise ValueError("Filtered HAMP workbook failed its row-count check")
    check.close()
    return output.getvalue()


def _style(sheet, table_name: str) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    if sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for cells in sheet.columns:
        letter = cells[0].column_letter
        width = max((len(str(cell.value or "")) for cell in cells[:200]), default=10)
        sheet.column_dimensions[letter].width = min(max(width + 2, 14), 48)
