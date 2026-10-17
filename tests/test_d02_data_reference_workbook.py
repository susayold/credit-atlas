import csv
from pathlib import Path

from openpyxl import load_workbook


def test_d02_reference_workbook_includes_each_referenced_public_data_csv():
    root = Path(__file__).parents[1]
    path = root / "apps/excel/templates/CreditAtlas-D02-Data-Reference.xlsx"
    assert path.is_file()
    workbook = load_workbook(path, read_only=False, data_only=True)
    assert workbook.sheetnames == [
        "Start here", "NSMO fields", "NSMO feature gates", "NSMO routing", "NSMO question refs", "NSMO categories", "NSMO wave coverage", "NSMO item wording",
        "NSMO returns", "NSMO usability", "NMDB quarterly", "NMDB states", "FHFA HPI state", "HUD monthly", "HUD subgroups", "HUD 90+ reasons",
        "HUD foreclosure", "HUD REO recovery", "HUD REO loss", "NMDB definitions", "HUD definitions", "NMDB state definitions", "FHFA HPI definitions",
    ]
    expected = {
        "NSMO fields": "nsmo-v70-field-dictionary.csv",
        "NSMO feature gates": "nsmo-v70-predictor-gates.csv",
        "NSMO routing": "nsmo-v70-routing-review.csv",
        "NSMO question refs": "nsmo-v70-question-reference-reconciliation.csv",
        "NSMO categories": "nsmo-v70-value-labels.csv",
        "NSMO wave coverage": "nsmo-v70-wave-coverage.csv",
        "NSMO item wording": "nsmo-v70-wave-item-comparability.csv",
        "NSMO returns": "nsmo-v70-survey-returns.csv",
        "NSMO usability": "nsmo-v70-usable-responses.csv",
        "NMDB quarterly": "nmdb-performance-national-quarterly-2026q2.csv",
        "NMDB states": "nmdb-performance-state-quarterly-2026q2.csv",
        "FHFA HPI state": "fhfa-state-all-transactions-hpi-quarterly-2026q2.csv",
        "HUD monthly": "hud-fha-performance-2026-08.csv",
        "HUD subgroups": "hud-fha-performance-2026-08-subgroups.csv",
        "HUD 90+ reasons": "hud-fha-new-90plus-reasons-2026-08.csv",
        "HUD foreclosure": "hud-fha-foreclosure-statistics-2026-08.csv",
        "HUD REO recovery": "hud-fha-reo-recovery-2026-08.csv",
        "HUD REO loss": "hud-fha-reo-loss-components-2026-08.csv",
    }
    available_csvs = {path.name for path in (root / "docs/data").glob("*.csv")}
    assert set(expected.values()).issubset(available_csvs)
    for sheet_name, csv_name in expected.items():
        with (root / "docs/data" / csv_name).open(encoding="utf-8-sig", newline="") as stream:
            row_count = sum(1 for _ in csv.DictReader(stream))
        sheet = workbook[sheet_name]
        assert sheet.max_row == row_count + 1
        assert sheet.freeze_panes == "B2"
        assert len(sheet.tables) == 1
        table = next(iter(sheet.tables.values()))
        assert table.autoFilter is not None
    assert workbook["NSMO categories"].cell(2, 2).value in {"1", "'1'"}
    assert workbook["NMDB quarterly"].max_row == 1569
    assert workbook["NMDB states"].max_row == 79969
    assert workbook["NMDB states"].cell(2, 12).value == "P3089DL"
    assert workbook["NMDB states"].cell(2, 13).value == 1.5
    assert "30–59 and 60–89" in workbook["NMDB definitions"].cell(2, 2).value
    assert "30–89 DPD" in workbook["NMDB definitions"].cell(2, 3).value
    assert workbook["NMDB state definitions"].max_row == 11
    assert "30–59 and 60–89" in workbook["NMDB state definitions"].cell(3, 2).value
    assert "30–89 DPD" in workbook["NMDB state definitions"].cell(3, 3).value
    assert "17 PFORB flags overlap" in workbook["NMDB state definitions"].cell(7, 3).value
    assert workbook["FHFA HPI state"].max_row == 10507
    hpi = workbook["FHFA HPI state"]
    assert hpi.cell(1, 4).value == "All-transactions HPI (NSA; 1980Q1=100)"
    assert hpi.cell(2, 1).value == "AK"
    assert hpi.cell(2, 2).value == 1975
    assert hpi.cell(2, 4).value == 69.15
    assert hpi.cell(10507, 1).value == "WY"
    assert hpi.cell(10507, 2).value == 2026
    assert hpi.cell(10507, 3).value == 2
    assert hpi.cell(10507, 4).value == 526.96
    assert workbook["FHFA HPI definitions"].max_row == 9
    assert "1980Q1=100" in workbook["FHFA HPI definitions"].cell(5, 2).value
    returns = workbook["NSMO returns"]
    assert returns.max_row == 49
    assert returns.cell(2, 1).value == 1
    assert returns.cell(47, 17).value is True
    assert returns.cell(48, 17).value is False
    assert returns.cell(49, 17).value is False
    usability = workbook["NSMO usability"]
    assert usability.max_row == 47
    assert usability.cell(2, 7).value == 4618
    assert usability.cell(8, 9).value == 295
    assert usability.cell(46, 9).value == 3
    assert usability.cell(47, 9).value == 494
    assert workbook["Start here"]["B2"].value.startswith("Excel copy of the D02")
    gates = workbook["NSMO feature gates"]
    assert gates.max_row == 588
    assert gates.cell(2, 2).value == "nsmoid"
    assert gates.cell(2, 7).value == "IDENTIFIER_EXCLUDE"
    assert all(not str(gates.cell(row, 9).value).startswith("APPROVED") for row in range(2, gates.max_row + 1))
    assert workbook["Start here"]["B4"].value.startswith("The feature-gates sheet")
    routing = workbook["NSMO routing"]
    assert routing.max_row == 320
    statuses = [routing.cell(row, 16).value for row in range(2, routing.max_row + 1)]
    ref_statuses = [routing.cell(row, 6).value for row in range(2, routing.max_row + 1)]
    assert statuses.count("CONDITIONAL_ROUTING_DOCUMENTED") == 58
    assert statuses.count("LEGACY_GATEWAY_MAPPING_REQUIRED") == 12
    assert ref_statuses.count("RECONCILED_FROM_FIELD_ID_AND_QUESTIONNAIRE") == 7
    question_refs = workbook["NSMO question refs"]
    assert question_refs.max_row == 8
    assert question_refs.cell(3, 1).value == "x17g_1"
    assert question_refs.cell(3, 4).value == "WAVE_VARIANT_REVIEW"
    assert workbook["HUD 90+ reasons"].max_row == 23
    assert workbook["HUD 90+ reasons"].cell(2, 1).value == "2021 Q4"
    assert workbook["HUD 90+ reasons"].cell(2, 13).value is False
    assert workbook["HUD foreclosure"].max_row == 38
    assert workbook["HUD foreclosure"].cell(38, 8).value == 0.28
    assert workbook["HUD REO recovery"].max_row == 51
    assert workbook["HUD REO recovery"].cell(2, 3).value == 50.5
    assert workbook["HUD REO loss"].max_row == 14
    assert workbook["HUD definitions"].max_row >= 9
