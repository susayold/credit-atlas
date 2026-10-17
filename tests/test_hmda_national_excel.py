from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from credit_atlas.hmda_national_excel import workbook_bytes


ACTION_NAMES = [
    "ORIGINATED", "APPROVED_NOT_ACCEPTED", "DENIED", "WITHDRAWN",
    "CLOSED_INCOMPLETE", "PURCHASED_LOAN", "PREAPPROVAL_DENIED",
    "PREAPPROVAL_APPROVED_NOT_ACCEPTED",
]


class FakeHmdaService:
    def __init__(self, *, selected=True):
        self.selected = selected
        self.hmda_national_views = self

    def report(self, view_id, query):
        result = {
            "view_id": view_id,
            "activity_year": 2025,
            "release_type": "static snapshot",
            "snapshot_date": "2026-06-02",
            "source_url": "https://ffiec.cfpb.gov/data-publication/snapshot-national-loan-level-dataset/2025",
            "limitations": ["Application decisions are not repayment performance."],
        }
        if self.selected:
            result["applied_filters"] = {"state_code": "CA", "loan_type": "1"}
            result["filtered_summary"] = {
                "records": 2,
                "decision_denominator": 2,
                "denied_numerator": 1,
                "descriptive_denial_rate": "0.5",
                "public_amount_sum_usd": "250000.00",
                "actions": [
                    {"code": str(index), "name": name, "records": 1 if index in (1, 3) else 0}
                    for index, name in enumerate(ACTION_NAMES, start=1)
                ],
            }
        return result

    def view(self, view_id):
        return {"view_id": view_id}


def test_filtered_hmda_workbook_contains_scope_kpis_actions_and_reconciliation_formulas():
    payload = workbook_bytes(FakeHmdaService(), {
        "view_id": "hmda-national-view-0123456789ab",
        "filter_state_code": "CA",
        "filter_loan_type": "1",
    })
    workbook = load_workbook(BytesIO(payload), data_only=False)

    assert workbook.sheetnames == ["Summary", "Actions", "Method & sources"]
    summary = workbook["Summary"]
    assert summary["B6"].value == "CA"
    assert summary["B7"].value == "Conventional · 1"
    assert summary["B12"].value == 2
    assert summary["B13"].value == 2
    assert summary["B14"].value == 1
    assert summary["B15"].value == 0.5
    assert summary["B19"].value.startswith("=IF(SUM(Actions!")
    assert summary["B22"].value.startswith("=IF(B13=0,")
    assert len(workbook["Actions"].tables) == 1
    assert workbook["Actions"]["D2"].value == "=C2/SUM($C$2:$C$9)"
    assert workbook["Method & sources"]["B5"].value == "HMDA action_taken in 1, 2, or 3 (originated, approved not accepted, or denied)."

    formulas = [cell.value for sheet in workbook.worksheets for row in sheet.iter_rows()
                for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert len(formulas) == 12
    workbook.close()


def test_filtered_hmda_workbook_requires_at_least_one_filter():
    with pytest.raises(ValueError, match="Select a state or loan type"):
        workbook_bytes(FakeHmdaService(selected=False), {"view_id": "hmda-national-view-0123456789ab"})
