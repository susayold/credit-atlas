import csv
from pathlib import Path

from openpyxl import load_workbook


def test_hud_fha_workbook_contains_full_curated_tables():
    root = Path(__file__).parents[1]
    workbook_path = root / "apps/excel/templates/CreditAtlas-HUD-FHA-Performance.xlsx"
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    assert workbook.sheetnames == ["Bat dau", "Monthly trend", "Subgroup rates"]
    expected = {
        "Monthly trend": "hud-fha-performance-2026-08.csv",
        "Subgroup rates": "hud-fha-performance-2026-08-subgroups.csv",
    }
    for sheet_name, csv_name in expected.items():
        with (root / "docs/data" / csv_name).open(encoding="utf-8-sig", newline="") as stream:
            expected_rows = sum(1 for _ in csv.DictReader(stream))
        sheet = workbook[sheet_name]
        assert sheet.max_row == expected_rows + 1
        assert sheet.freeze_panes == ("D2" if sheet_name == "Subgroup rates" else "A2")
        assert len(sheet.tables) == 1
        assert next(iter(sheet.tables.values())).autoFilter is not None
    assert workbook["Monthly trend"][14][0].value == "2026-08"
    assert workbook["Subgroup rates"].max_row == 61
    assert workbook["Subgroup rates"][14][12].value == 11.28
    assert "HTTP 403" in workbook["Bat dau"][8][1].value
