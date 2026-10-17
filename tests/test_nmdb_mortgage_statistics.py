import json
import io
from pathlib import Path

import duckdb
import pytest
from openpyxl import load_workbook

from credit_atlas.nmdb_mortgage_statistics import MART_RELATIVE, RECEIPT_RELATIVE, report
from credit_atlas.nmdb_mortgage_statistics_excel import workbook_bytes
from credit_atlas.workspace import sha256


NEW_VALUE1_ONLY = ("TOT_ORIG", "TOT_LOANAMT", "AVE_LOANAMT", "AVE_PAYMENT", "AVE_PROPVAL")


@pytest.fixture
def nmdb_root(tmp_path):
    root = tmp_path
    mart_path = root / MART_RELATIVE
    mart_path.parent.mkdir(parents=True)
    rows = [
        ("new_annual", "NMDB", "Annual", "National", "USA", "United States", "All Mortgages",
         "2025", 2025, None, None, False, series, str(i + 1), str(i + 11), float(i + 1), float(i + 11))
        for i, series in enumerate((*NEW_VALUE1_ONLY, "AVE_INTRATE"))
    ]
    con = duckdb.connect()
    try:
        con.execute("""
            CREATE TABLE mart (
                product VARCHAR, source VARCHAR, frequency VARCHAR, geography_level VARCHAR,
                geography_id VARCHAR, geography_name VARCHAR, market VARCHAR, period VARCHAR,
                year INTEGER, quarter INTEGER, month INTEGER, suppressed BOOLEAN, series_id VARCHAR,
                value1_raw VARCHAR, value2_raw VARCHAR, value1 DOUBLE, value2 DOUBLE
            )
        """)
        con.executemany("INSERT INTO mart VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", rows)
        con.execute("COPY mart TO ? (FORMAT PARQUET)", [str(mart_path)])
    finally:
        con.close()

    source_receipt = root / "source-receipt.json"
    source_receipt.write_text(json.dumps({"status": "PASS"}), encoding="utf-8")
    receipt_path = root / RECEIPT_RELATIVE
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    receipt = {
        "status": "PASS",
        "source_release_date": "2026-09-30",
        "retrieved_at": "2026-10-16T14:20:09Z",
        "row_count": len(rows),
        "source_evidence": {"new_annual": {"source_rows": len(rows)}},
        "source_receipt": str(source_receipt),
        "source_receipt_sha256": sha256(source_receipt),
        "mart_sha256": sha256(mart_path),
        "quality": {},
    }
    receipt_path.write_text(json.dumps(receipt), encoding="utf-8")
    return root


@pytest.mark.parametrize("series_id", NEW_VALUE1_ONLY)
def test_appendix_b_value1_only_rule_blocks_raw_value2_even_when_present(nmdb_root, series_id):
    with pytest.raises(ValueError, match="Selected weighting has no observed values"):
        report(nmdb_root, {
            "product": "new_annual", "geo_level": "National", "geo_id": "USA",
            "market": "All Mortgages", "series_id": series_id, "measure": "VALUE2",
            "from_year": "2025", "to_year": "2025",
        })


def test_value2_remains_selectable_for_a_series_where_appendix_b_supports_it(nmdb_root):
    result = report(nmdb_root, {
        "product": "new_annual", "geo_level": "National", "geo_id": "USA",
        "market": "All Mortgages", "series_id": "AVE_INTRATE", "measure": "VALUE2",
        "from_year": "2025", "to_year": "2025",
    })

    assert result["rows"][0]["selected_value"] == 16.0
    assert result["measure"] == "VALUE2"
    assert result["quality"]["forecast_or_imputation_used"] is False


def test_nmdb_market_workbook_uses_table_filter_without_duplicate_sheet_filter(nmdb_root):
    query = {
        "product": "new_annual", "geo_level": "National", "geo_id": "USA",
        "market": "All Mortgages", "series_id": "AVE_INTRATE", "measure": "VALUE2",
        "from_year": "2025", "to_year": "2025",
    }

    class Service:
        @staticmethod
        def nmdb_mortgage_statistics(filters):
            return report(nmdb_root, filters)

    workbook = load_workbook(io.BytesIO(workbook_bytes(Service(), query)), data_only=False)
    data = workbook["Selected data"]

    assert data.auto_filter.ref is None
    table = data.tables["NMDBSelectedData"]
    assert table.ref == "A1:R2"
    assert table.autoFilter.ref == "A1:R2"
    assert data.cell(2, 13).value == 16.0
