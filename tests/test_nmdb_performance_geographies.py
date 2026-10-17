import csv
import io
import zipfile

import pytest
from openpyxl import load_workbook

from credit_atlas import nmdb_performance_geographies as module


def _row(level, geo_id, name, period, series, value, *, suppressed="0"):
    year, quarter = period.split("Q")
    return {
        "SOURCE": "NMDB", "FREQUENCY": "Quarterly", "GEOLEVEL": level,
        "GEOID": geo_id, "GEONAME": name, "MARKET": "All Mortgages",
        "PERIOD": period, "YEAR": year, "QUARTER": quarter, "MONTH": "06",
        "SUPPRESSED": suppressed, "SERIESID": series, "VALUE1": value,
    }


@pytest.fixture
def receipt_profiled_package(tmp_path, monkeypatch):
    rows = [
        _row("National", "USA", "United States", "2002Q1", "PFORB", ""),
        _row("State", "CA", "California", "2026Q2", "P90DL", "1.1"),
        _row("State", "VT", "Vermont", "2004Q1", "P90DL", "", suppressed="1"),
    ]
    csv_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(csv_buffer, fieldnames=module.EXPECTED_COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    archive_path = tmp_path / "performance.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("performance.csv", csv_buffer.getvalue())
    profile = {
        "member": "performance.csv", "rows": 3, "unique_aggregate_keys": 3,
        "geography_level_row_counts": {"National": 1, "State": 2},
        "market_row_counts": {"All Mortgages": 3},
        "series_row_counts": {"PFORB": 1, "P90DL": 2},
        "suppression_row_counts": {"0": 2, "1": 1},
        "blank_or_dot_value_counts": {"VALUE1": 2},
        "csv_sha256": "test-csv-sha",
    }
    source = {
        "status": "AVAILABLE", "available": True, "source_url": module.SOURCE_URL,
        "source_page": module.INDEX_URL, "source_release_date": "2026-Q2",
        "retrieved_at": "2026-10-17T00:00:00Z", "dataset": "FHFA NMDB test",
        "quality": {"archive_sha256": "test-archive-sha"},
    }
    monkeypatch.setattr(module, "aggregate_read", lambda _root: source)
    monkeypatch.setattr(module, "_full_package", lambda _root, _sha: (archive_path, profile))
    return tmp_path


def test_nmdb_geography_preserves_suppressed_and_pre_series_blanks(receipt_profiled_package):
    report = module.report(receipt_profiled_package, {
        "geo_level": "State", "geo_id": "VT", "market": "All Mortgages",
        "series_id": "P90DL", "period": "2004Q1",
    })

    assert report["coverage"]["selected_rows"] == 1
    assert report["coverage"]["unavailable_values"] == 1
    assert report["coverage"]["suppressed_values"] == 1
    assert report["rows"][0]["value_pct"] is None
    assert report["rows"][0]["value_state"] == "suppressed"

    early_forbearance = module.report(receipt_profiled_package, {
        "geo_level": "National", "geo_id": "USA", "series_id": "PFORB", "period": "2002Q1",
    })
    assert early_forbearance["rows"][0]["value_pct"] is None
    assert early_forbearance["rows"][0]["value_state"] == "unavailable_before_series_start"


def test_nmdb_geography_workbook_contains_query_lineage_and_keeps_blank(receipt_profiled_package):
    report = module.report(receipt_profiled_package, {
        "geo_level": "National", "geo_id": "USA", "series_id": "PFORB", "period": "2002Q1",
    })
    workbook = load_workbook(io.BytesIO(module.workbook_bytes(report)), data_only=True)

    assert workbook.sheetnames == ["Start Here", "Selected data", "Definitions", "Quality & lineage"]
    assert workbook["Selected data"].cell(2, 10).value is None
    assert workbook["Selected data"].cell(2, 12).value == "unavailable_before_series_start"
    assert "test-archive-sha" in workbook["Quality & lineage"][9][1].value


def test_public_national_fallback_does_not_claim_full_geography_data(tmp_path, monkeypatch):
    monkeypatch.setattr(module, "aggregate_read", lambda _root: {
        "status": "AVAILABLE", "available": True, "source_url": "https://www.fhfa.gov/national-only.zip",
        "source_page": module.INDEX_URL, "quality": {"archive_sha256": "national-only"},
    })

    report = module.report(tmp_path)

    assert report["available"] is False
    assert "national fallback only" in report["reason"]


@pytest.mark.parametrize("query", [
    {"geo_level": "County"},
    {"series_id": "DEFAULT"},
    {"period": "2026Q3"},
    {"unexpected": "value"},
])
def test_nmdb_geography_rejects_unpublished_filter_values(query):
    with pytest.raises(ValueError):
        module._filters(query)
