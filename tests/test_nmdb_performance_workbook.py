import csv
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts import add_nmdb_state_to_reference_workbook as state_workbook_builder


def test_nmdb_workbook_has_latest_summary_and_filterable_full_series():
    root = Path(__file__).parents[1]
    workbook = load_workbook(root / "apps/excel/templates/CreditAtlas-NMDB-Performance.xlsx", data_only=True)
    assert workbook.sheetnames == ["Bat dau", "Latest KPIs", "Quarterly data", "Definitions"]
    with (root / "docs/data/nmdb-performance-national-quarterly-2026q2.csv").open(encoding="utf-8-sig", newline="") as stream:
        source_rows = list(csv.DictReader(stream))
    latest = workbook["Latest KPIs"]
    assert latest.max_row == 5
    assert latest["A2"].value == "2026Q2"
    assert latest["D1"].value == "30–89 DPD (%)"
    assert latest["D2"].value == 1.7
    assert latest["E2"].value == 1.0
    assert latest["G2"].value == 0.5
    assert latest.freeze_panes == "B2"
    assert len(latest.tables) == 1
    quarterly = workbook["Quarterly data"]
    assert quarterly.max_row == len(source_rows) + 1
    assert quarterly["G2"].value == "2002Q1"
    assert quarterly["G1569"].value == "2026Q2"
    assert quarterly["M4"].value is None
    assert quarterly.freeze_panes == "B2"
    assert len(quarterly.tables) == 1
    assert next(iter(quarterly.tables.values())).autoFilter is not None
    assert "2019Q4" in workbook["Definitions"]["D5"].value
    assert workbook["Definitions"]["B2"].value == "30–89 days delinquent"
    assert "30–59 or 60–89" in workbook["Definitions"]["C2"].value


def test_nmdb_state_reference_table_has_filter_without_duplicate_sheet_filter(tmp_path, monkeypatch):
    data_path = tmp_path / "state-panel.csv"
    data_path.write_text(
        "source,frequency,geo_level,geo_id,geo_name,market,period,year,quarter,month,suppressed,series_id,value_pct\n"
        "NMDB,Quarterly,State,VT,Vermont,All Mortgages,2004Q1,2004,1,3,1,P3089DL,\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(state_workbook_builder, "DATA", data_path)
    workbook = Workbook()
    workbook.active.title = "Start here"
    workbook.create_sheet("FHFA HPI state")

    state_workbook_builder.add_data_sheet(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    states = load_workbook(output)["NMDB states"]

    assert states.auto_filter.ref is None
    table = states.tables["NmdbStates"]
    assert table.ref == "A1:M2"
    assert table.autoFilter.ref == "A1:M2"
    assert states.cell(2, 13).value is None
