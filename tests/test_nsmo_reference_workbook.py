import csv
from pathlib import Path

from openpyxl import load_workbook

def test_nsmo_reference_workbook_contains_all_reviewed_metadata():
    root = Path(__file__).parents[1]
    output = root / "apps/excel/templates/CreditAtlas-NSMO-Reference.xlsx"
    assert output.is_file()
    workbook = load_workbook(output, read_only=False, data_only=True)
    assert workbook.sheetnames == [
        "Bat dau", "Tu dien 587 truong", "Nhan va category", "Coverage theo wave", "So sanh wording item"
    ]
    expected = {
        "Tu dien 587 truong": "nsmo-v70-field-dictionary.csv",
        "Nhan va category": "nsmo-v70-value-labels.csv",
        "Coverage theo wave": "nsmo-v70-wave-coverage.csv",
        "So sanh wording item": "nsmo-v70-wave-item-comparability.csv",
    }
    for sheet_name, csv_name in expected.items():
        with (root / "docs" / "data" / csv_name).open(encoding="utf-8-sig", newline="") as stream:
            row_count = sum(1 for _ in csv.DictReader(stream))
        sheet = workbook[sheet_name]
        assert sheet.max_row == row_count + 1
        # The workbook freezes both the header row and the identifier column.
        assert sheet.freeze_panes == "B2"
        assert len(sheet.tables) == 1
        table = next(iter(sheet.tables.values()))
        assert table.ref.startswith("A1:")
        assert table.autoFilter is not None
    labels = workbook["Nhan va category"]
    assert labels[1][1].value == "Mã category (giữ nguyên)"
    assert labels.max_row > 100
    assert workbook["Bat dau"]["B4"].value.startswith("Bốn sheet dữ liệu")
