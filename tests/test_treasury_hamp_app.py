from __future__ import annotations

import csv
import json
import threading
from io import BytesIO
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from urllib.error import HTTPError
from urllib.request import urlopen

import pytest
from openpyxl import load_workbook

from credit_atlas.server import handler_for
from credit_atlas.treasury_hamp_quarterly import CSV_RELATIVE_PATH, report, workbook_bytes


CSV_FIELDS = [
    "source_id", "publication_quarter", "updated_through_date", "program_segment",
    "modification_effective_year", "permanent_modifications_started_as_reported",
    "modification_age_months", "delinquency_60_plus_pct", "delinquency_90_plus_pct",
    "cell_status", "source_pdf", "source_url", "source_pdf_sha256",
]


def _workspace(root: Path) -> None:
    path = root / CSV_RELATIVE_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    records = [
        ["hamp-2022q4", "2022-Q4", "2022-12-31", "HAMP Tier 1", "2009", "63359", "12", "20.1", "15.2", "reported_in_source", "private.pdf", "https://home.treasury.gov/2022q4.pdf", "a" * 64],
        ["hamp-2023q1", "2023-Q1", "2023-03-31", "HAMP Tier 1", "2009", "63359", "12", "", "", "not_available_in_published_table", "private.pdf", "https://home.treasury.gov/2023q1.pdf", "b" * 64],
        ["hamp-2023q1", "2023-Q1", "2023-03-31", "HAMP Tier 1", "2010", "54000", "12", "22.0", "17.1", "reported_in_source", "private.pdf", "https://home.treasury.gov/2023q1.pdf", "b" * 64],
    ]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(CSV_FIELDS)
        writer.writerows(records)


def test_hamp_report_filters_exact_cohort_and_preserves_blank(tmp_path):
    _workspace(tmp_path)
    result = report(tmp_path, {
        "publication_quarter": "all", "program_segment": "HAMP Tier 1",
        "modification_effective_year": "2009", "modification_age_months": "12",
    })
    assert result["coverage"]["source_reports"] == 2
    assert result["quality"] == {
        "selected_cells": 2, "reported_cells": 1, "unavailable_cells": 1,
        "distinct_cohorts_across_snapshots": 2, "source_reports": 2,
    }
    assert result["rows"][-1]["cell_status"] == "not_available_in_published_table"
    assert result["rows"][-1]["delinquency_60_plus_pct"] is None
    assert result["rows"][-1]["delinquency_90_plus_pct"] is None
    assert "source_pdf" not in result["rows"][-1]


def test_hamp_report_rejects_unknown_filters_and_does_not_pool(tmp_path):
    _workspace(tmp_path)
    with pytest.raises(ValueError, match="filter value"):
        report(tmp_path, {"publication_quarter": "2021-Q3"})
    all_rows = report(tmp_path)
    assert all_rows["quality"]["selected_cells"] == 3
    assert all_rows["rows"][0]["program_segment"] == "HAMP Tier 1"
    assert all_rows["quality"]["distinct_cohorts_across_snapshots"] == 3


def test_missing_private_hamp_data_is_explicit_and_has_no_sample_rows(tmp_path):
    result = report(tmp_path)
    assert result["available"] is False
    assert result["rows"] == []
    assert "not installed" in result["reason"]


def test_hamp_filtered_workbook_is_excel_readable_and_keeps_missing_blank(tmp_path):
    _workspace(tmp_path)
    result = report(tmp_path, {"publication_quarter": "2023-Q1", "program_segment": "HAMP Tier 1",
                               "modification_effective_year": "2009", "modification_age_months": "12"})
    book = load_workbook(BytesIO(workbook_bytes(result)), data_only=True)
    assert book.sheetnames == ["Cohort observations", "Source index", "Method & limits"]
    sheet = book["Cohort observations"]
    assert sheet.max_row == 2
    assert sheet["G2"].value is None and sheet["H2"].value is None
    assert "not_available_in_published_table" in sheet["I2"].value
    assert "2023q1.pdf" in sheet["K2"].value
    book.close()


def test_hamp_local_api_and_filtered_excel_export(tmp_path):
    _workspace(tmp_path)
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler_for(SimpleNamespace(root=tmp_path)))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        with urlopen(base + "/api/treasury-hamp-quarterly?program_segment=HAMP+Tier+1&modification_effective_year=2009&modification_age_months=12") as response:
            payload = json.loads(response.read())
            assert response.status == 200
            assert payload["quality"]["selected_cells"] == 2
        with urlopen(base + "/api/treasury-hamp-quarterly-excel?publication_quarter=2023-Q1&program_segment=HAMP+Tier+1&modification_effective_year=2009&modification_age_months=12") as response:
            body = response.read()
            assert response.status == 200
            assert response.headers.get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            assert body[:2] == b"PK"
        with pytest.raises(HTTPError) as invalid:
            urlopen(base + "/api/treasury-hamp-quarterly?publication_quarter=2021-Q3")
        assert invalid.value.code == 400
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
