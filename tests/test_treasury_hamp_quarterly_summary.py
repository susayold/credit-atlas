from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from scripts.build_treasury_hamp_quarterly_summary import CurationError, parse_summary_text, style_table
from scripts.download_treasury_hamp_quarterly_summaries import load_sources


REPORT = {
    "source_id": "treasury-hamp-summary-2023q1",
    "quarter": "2023-Q1",
    "updated_through": "2023-03-31",
    "file": "summary.pdf",
    "url": "https://home.treasury.gov/example.pdf",
}


def test_quarterly_source_index_is_complete_and_official_host_only():
    _, reports = load_sources()
    assert len(reports) == 20
    assert len({row["source_id"] for row in reports}) == 20
    assert "2021-Q3" not in {row["quarter"] for row in reports}
    assert all(row["url"].startswith("https://home.treasury.gov/") for row in reports)
    assert all("published_on" in row and "updated_through" in row for row in reports)


def sample_layout(tier1_last_row: str = "") -> str:
    full = "10.0% 5.0% 20.0% 10.0%"
    lines = ["Performance of HAMP Permanent Modifications - Tier 1"]
    lines.extend(f"{year} 1,000 {full}" for year in range(2009, 2018))
    lines.append("Performance of HAMP Permanent Modifications - Tier 2")
    lines.extend(f"{year} 100 {full}" for year in range(2012, 2018))
    lines.append("Performance of HAMP Permanent Modifications - Streamline HAMP")
    lines.extend(f"{year} 50 {full}" for year in (2016, 2017))
    if tier1_last_row:
        lines.insert(10, tier1_last_row)
    return "\n".join(lines)


def test_quarterly_summary_parser_preserves_trailing_unavailable_horizons():
    rows = parse_summary_text(sample_layout(), "Updated through March 31, 2023", REPORT, "a" * 64)
    assert len(rows) == 17 * 7
    first = rows[0]
    assert first["program_segment"] == "HAMP Tier 1"
    assert first["modification_effective_year"] == "2009"
    assert first["modification_age_months"] == 6
    assert first["delinquency_60_plus_pct"] == 10.0
    assert first["delinquency_90_plus_pct"] == 5.0
    third_age = next(row for row in rows if row["program_segment"] == "HAMP Tier 1"
                     and row["modification_effective_year"] == "2009"
                     and row["modification_age_months"] == 24)
    assert third_age["delinquency_60_plus_pct"] is None
    assert third_age["cell_status"] == "not_available_in_published_table"


def test_quarterly_summary_parser_rejects_wrong_cutoff():
    with pytest.raises(CurationError, match="cutoff"):
        parse_summary_text(sample_layout(), "Updated through December 31, 2022", REPORT, "a" * 64)


def test_quarterly_summary_parser_rejects_odd_metric_pairs():
    layout = sample_layout(tier1_last_row="2009 1,000 10.0%")
    with pytest.raises(CurationError, match="Malformed percentage row"):
        parse_summary_text(layout, "Updated through March 31, 2023", REPORT, "a" * 64)


def test_quarterly_summary_parser_rejects_invalid_90_plus_subset():
    layout = sample_layout().replace("2009 1,000 10.0% 5.0%", "2009 1,000 10.0% 11.0%", 1)
    with pytest.raises(CurationError, match="subset/range"):
        parse_summary_text(layout, "Updated through March 31, 2023", REPORT, "a" * 64)


def test_treasury_summary_table_uses_table_filter_without_duplicate_sheet_filter():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cohort KPIs"
    sheet.append(["source_id", "value"])
    sheet.append(["hamp-2023q1", 1.25])

    style_table(sheet, "HampQuarterlyCohortKpis")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    sheet = load_workbook(output)["Cohort KPIs"]

    assert sheet.auto_filter.ref is None
    table = sheet.tables["HampQuarterlyCohortKpis"]
    assert table.ref == "A1:B2"
    assert table.autoFilter.ref == "A1:B2"
